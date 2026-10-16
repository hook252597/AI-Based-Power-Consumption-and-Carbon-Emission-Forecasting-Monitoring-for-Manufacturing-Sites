# -*- coding: utf-8 -*-
"""마가렛트 에이전트 — 도구 함수 12개 (agent/agent.md 1절·5절)

필요 패키지 : pandas, pyarrow, openpyxl, python-docx(.docx 양식 파싱)
환경변수    : 없음 (이 모듈은 API 키를 읽지 않는다)
실행        : 단독 실행하지 않는다. agent_core.run_turn / agent_ui.render_agent_tab 이 사용한다.

데이터(읽기 전용) — 대시보드와 같은 산출물을 쓴다 (tools/prepare_dashboard_cache.py):
  ml/outputs/agent_data/dashboard_hourly.parquet    2024-12-01~2025-04-29 설비·시간별 실측 + AI 판정
  ml/outputs/agent_data/dashboard_forecast.parquet  2025-04-30~2025-05-31 설비·시간별 AI 전력 예측
  ml/outputs/agent_data/dashboard_hourly.json       설비별 AI 경고 기준(ai_thresholds)
  ml/outputs/forecast_v2/summary.json               예측 모델 이름·MAE
  (동업자 에이전트 agent_runtime 의 코드·캐시는 사용하지 않는다)

보고서 양식:
  agent/BasicSafetyReportForm.xlsx 기본 양식은 읽기 전용이다.
  변경 양식·업로드 원본·초안은 현재 브라우저 세션에만 보관한다.
  종합판정은 담당자의 화면 확인 뒤에만 내보낼 수 있다.

모든 도구는 JSON 직렬화 가능한 dict/list를 돌려주고, 실패해도 예외를 밖으로 던지지 않는다.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any

import pandas as pd

# -----------------------------------------------------------------------------
# 상수 — 대시보드(dashboard_app.py)와 같은 값
# -----------------------------------------------------------------------------
APP_DIR = Path(__file__).resolve().parent
REPO = APP_DIR.parent
ML = REPO / "ml"
HOURLY_PATH = ML / "outputs/agent_data/dashboard_hourly.parquet"
FORECAST_PATH = ML / "outputs/agent_data/dashboard_forecast.parquet"
SUMMARY_PATH = ML / "outputs/forecast_v2/summary.json"
HOURLY_META_PATH = ML / "outputs/agent_data/dashboard_hourly.json"   # ai_thresholds 포함(같은 스크립트가 생성)

NOW = datetime(2025, 4, 30, 0, 0)   # 고정 기준 시점(대시보드와 동일)
TODAY = (NOW - timedelta(seconds=1)).date()
WON_PER_KWH = 180.0
KGCO2_PER_KWH = 0.424
MODULE = "module(equipment)"
DEFAULT_THRESHOLDS = {"phase_voltage": 210.0, "line_voltage": 360.0, "power_factor": 85.0}

# 설비 13개 — 앞의 숫자는 RTU 채널 번호(다시 매기지 않는다)
EQUIPMENT_LIST = [
    ("1", "1 (PM-3)"), ("2", "2 (L-1전등)"), ("3", "3 (분쇄기(2))"), ("4", "4 (분쇄기(1))"),
    ("5", "5 (좌측분전반)"), ("11", "11 (우측분전반1)"), ("12", "12 (4호기)"), ("13", "13 (3호기)"),
    ("14", "14 (2호기)"), ("15", "15 (예비건조기)"), ("16", "16 (호이스트)"), ("17", "17 (6호기)"),
    ("18", "18 (우측분전반2)"),
]
NAME_BY_CHANNEL = dict(EQUIPMENT_LIST)

# 측정 항목: API 이름 → (컬럼, 표시명, 단위, 임계치 키)
ITEMS = {
    "phase_voltage_R": ("voltageR", "상전압 R", "V", "phase_voltage"),
    "phase_voltage_S": ("voltageS", "상전압 S", "V", "phase_voltage"),
    "phase_voltage_T": ("voltageT", "상전압 T", "V", "phase_voltage"),
    "line_voltage_RS": ("voltageRS", "선간전압 RS", "V", "line_voltage"),
    "line_voltage_ST": ("voltageST", "선간전압 ST", "V", "line_voltage"),
    "line_voltage_TR": ("voltageTR", "선간전압 TR", "V", "line_voltage"),
    "power_factor_R": ("powerFactorR", "역률 R", "%", "power_factor"),
    "power_factor_S": ("powerFactorS", "역률 S", "%", "power_factor"),
    "power_factor_T": ("powerFactorT", "역률 T", "%", "power_factor"),
}
FAMILY_LABEL = {"phase_voltage": "상전압", "line_voltage": "선간전압", "power_factor": "역률"}

DEFAULT_TEMPLATE_FILE = "BasicSafetyReportForm.xlsx"
DEFAULT_TEMPLATE_NAME = "제조현장 안전 점검 보고서"
AUTO_FILLS = {"FFEAF3F2"}   # 연한 청록 → 데이터가 자동으로 채움
USER_FILLS = {"FFFFF9E0"}   # 연한 노랑 → 사람이 적음
BLANK_NAME = re.compile(r"문서\s*번호|결재|서명|직인|도장|승인\s*란")
DEFAULT_REQUIRED = {"점검일시", "점검구분", "점검자", "종합판정", "설비명"}
PARSER_VERSION = 2


def _norm(text: Any) -> str:
    return re.sub(r"[\s_·]+", "", str(text or "")).lower()


def _iso(d: date) -> str:
    return d.isoformat()


def _parse_date(value: Any) -> date | None:
    try:
        return pd.Timestamp(str(value)[:10]).date()
    except (ValueError, TypeError):
        return None


def resolve_equipment(value: Any) -> str | None:
    """'12 (4호기)', '12(4호기)', '4호기', '12', 'PM-3' → 대시보드 표기 '12 (4호기)'. None/전체 → None.
    등록되지 않은 이름이면 ValueError."""
    if value is None or _norm(value) in ("", "전체", "공장전체", "all", "none", "null"):
        return None
    token = re.sub(r"[\s()（）\-]+", "", str(value)).lower()
    for ch, name in EQUIPMENT_LIST:
        label = name.split("(", 1)[1].rsplit(")", 1)[0]
        if token in (re.sub(r"[\s()（）\-]+", "", name).lower(), re.sub(r"[\s()（）\-]+", "", label).lower()):
            return name
    m = re.fullmatch(r"(?:ch|채널)?(\d+)(?:번|채널)?", token)
    if m and m.group(1) in NAME_BY_CHANNEL:
        return NAME_BY_CHANNEL[m.group(1)]
    raise ValueError(f"등록되지 않은 설비입니다: {value}")


# -----------------------------------------------------------------------------
# 데이터 접근 (파일 변경 시각을 키로 캐시)
# -----------------------------------------------------------------------------
def _mtime(path: Path) -> float:
    return path.stat().st_mtime if path.is_file() else 0.0


@lru_cache(maxsize=2)
def _hourly(mtime: float) -> pd.DataFrame:
    if not mtime:
        return pd.DataFrame()
    df = pd.read_parquet(HOURLY_PATH)
    df["channel"] = df[MODULE].str.split("(", n=1).str[0]
    df["equipment"] = df["channel"].map(NAME_BY_CHANNEL)
    df = df.dropna(subset=["equipment"])
    df = df[df["datetime"] + pd.Timedelta(hours=1) <= pd.Timestamp(NOW)]  # 끝난 시간만
    df["date"] = df["datetime"].dt.date
    df["hour"] = df["datetime"].dt.hour
    return df.reset_index(drop=True)


@lru_cache(maxsize=2)
def _forecast(mtime: float) -> pd.DataFrame:
    if not mtime:
        return pd.DataFrame()
    df = pd.read_parquet(FORECAST_PATH)
    df["channel"] = df[MODULE].str.split("(", n=1).str[0]
    df["equipment"] = df["channel"].map(NAME_BY_CHANNEL)
    df = df.dropna(subset=["equipment"])
    df = df[df["datetime"] >= pd.Timestamp(NOW)]
    df["date"] = df["datetime"].dt.date
    df["hour"] = df["datetime"].dt.hour
    return df.rename(columns={"prediction_kWh": "kwh"}).reset_index(drop=True)


def hourly() -> pd.DataFrame:
    return _hourly(_mtime(HOURLY_PATH))


def forecast() -> pd.DataFrame:
    return _forecast(_mtime(FORECAST_PATH))


@lru_cache(maxsize=1)
def _model_meta() -> dict:
    try:
        summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
        names = {"extra_trees": "Extra Trees", "random_forest": "Random Forest", "xgboost": "XGBoost",
                 "hist_gradient_boosting": "HistGradientBoosting"}
        return {"model": names.get(summary.get("selected_family"), summary.get("selected_family")),
                "mae_kwh_per_hour": float(summary["selected_may"]["mae_kWh_if_reference_is_Wh"])}
    except (OSError, ValueError, KeyError, TypeError):
        return {"model": None, "mae_kwh_per_hour": None}


@lru_cache(maxsize=2)
def _ai_thresholds(mtime: float) -> dict:
    if not mtime:
        return {}
    try:
        raw = json.loads(HOURLY_META_PATH.read_text(encoding="utf-8")).get("ai_thresholds") or {}
    except (OSError, ValueError):
        return {}
    out = {}
    for module, value in raw.items():
        ch = str(module).split("(", 1)[0]
        if ch in NAME_BY_CHANNEL and value is not None:
            out[NAME_BY_CHANNEL[ch]] = float(value)
    return out


def ai_thresholds() -> dict:
    """설비별 AI 이상점수 경고 기준(대시보드 표기 → 값). dashboard_hourly.json에 없으면 빈 dict.
    채우는 방법: python tools/prepare_dashboard_cache.py --thresholds-only"""
    return _ai_thresholds(_mtime(HOURLY_META_PATH))


def sensor_period() -> tuple[str, str] | None:
    h = hourly()
    if h.empty:
        return None
    return h["datetime"].min().strftime("%Y-%m-%d"), (h["datetime"].max() + pd.Timedelta(hours=1)).strftime("%Y-%m-%d %H:%M")


def forecast_period() -> tuple[str, str] | None:
    f = forecast()
    if f.empty:
        return None
    return f["datetime"].min().strftime("%Y-%m-%d"), f["datetime"].max().strftime("%Y-%m-%d %H:%M")


def _day(d: date, equipment: str | None = None) -> pd.DataFrame:
    h = hourly()
    if h.empty:
        return h
    out = h[h["date"] == d]
    return out[out["equipment"] == equipment] if equipment else out


def _no_data_note(d: date) -> str:
    period = sensor_period()
    span = f"{period[0]}~{period[1]}" if period else "없음"
    if d >= NOW.date():
        fp = forecast_period()
        hint = f" 전력·요금은 AI 예측({fp[0]}~{fp[1][:10]})으로만 확인할 수 있습니다." if fp and fp[0] <= _iso(d) <= fp[1][:10] else ""
        return f"{d}은(는) 기준 시점({NOW:%Y-%m-%d %H:%M}) 이후라 실측이 없습니다.{hint}"
    return f"{d} 의 설비 실측 자료가 없습니다. 실측 조회 가능 기간: {span}."


def _month_hourly(year: int, month: int, equipment: str | None = None) -> pd.DataFrame:
    """월의 시간별 전력량: 실측 우선, 없으면 AI 예측. 컬럼: datetime, actual, pred, value"""
    start = pd.Timestamp(year, month, 1)
    idx = pd.date_range(start, start + pd.offsets.MonthBegin(1), freq="h", inclusive="left")
    h, f = hourly(), forecast()
    if equipment:
        h = h[h["equipment"] == equipment] if len(h) else h
        f = f[f["equipment"] == equipment] if len(f) else f
    n = 1 if equipment else len(EQUIPMENT_LIST)
    if len(h):
        g = h.groupby("datetime").agg(kwh=("kwh", "sum"), m=("equipment", "nunique"), r=("n", "sum"))
        actual = g.loc[(g["m"] == n) & (g["r"] == 720 * n), "kwh"]
    else:
        actual = pd.Series(dtype=float)
    pred = f.groupby("datetime")["kwh"].sum() if len(f) else pd.Series(dtype=float)
    out = pd.DataFrame({"datetime": idx, "actual": actual.reindex(idx).to_numpy(), "pred": pred.reindex(idx).to_numpy()})
    out["value"] = out["actual"].where(out["actual"].notna(), out["pred"])
    return out


def _month_estimate(year: int, month: int, equipment: str | None = None) -> dict | None:
    hv = _month_hourly(year, month, equipment)
    valid = hv["value"].dropna()
    if valid.empty:
        return None
    total = float(valid.sum() + valid.mean() * (len(hv) - len(valid)))
    return {"total": total, "actual": float(hv["actual"].sum(skipna=True)),
            "actual_hours": int(hv["actual"].notna().sum()), "pred": float(hv.loc[hv["actual"].isna(), "pred"].sum(skipna=True)),
            "pred_hours": int((hv["actual"].isna() & hv["pred"].notna()).sum()),
            "filled_hours": int(hv["value"].isna().sum()), "hours": len(hv)}


def _money(kwh: float | None) -> dict:
    if kwh is None:
        return {"kwh": None, "cost_krw": None, "co2_kg": None}
    return {"kwh": round(kwh, 3), "cost_krw": round(kwh * WON_PER_KWH), "co2_kg": round(kwh * KGCO2_PER_KWH, 3)}


# -----------------------------------------------------------------------------
# 이상 판정 (시간 단위 집계: 해당 시간 안의 5초 측정 최저값 기준)
# -----------------------------------------------------------------------------
def alert_summary(d: date, thresholds: dict, equipment: str | None = None) -> dict:
    from monitoring.alert_service import get_alert_summary
    return get_alert_summary(_iso(d), thresholds=thresholds, equipment=equipment, cutoff=NOW)


def _events(d: date, thresholds: dict, equipment: str | None = None) -> tuple[list[dict], list[dict]]:
    summary = alert_summary(d, thresholds, equipment)
    return summary.get("rule_events", []), summary.get("ai_events", [])


def duration_text(event: dict) -> str:
    seconds = event.get("duration_seconds")
    if seconds is None:
        return "확인 불가 · 이상 기록이 관측된 시간대"
    seconds = int(seconds)
    if seconds < 60:
        return f"{seconds}초"
    minutes, seconds = divmod(seconds, 60)
    return f"{minutes}분" + (f" {seconds}초" if seconds else "")


def _status(d: date, thresholds: dict) -> dict[str, dict]:
    summary = alert_summary(d, thresholds)
    rows = {row["equipment"]: row for row in summary.get("equipment_rows", [])}
    out = {}
    for _, name in EQUIPMENT_LIST:
        row = rows.get(name, {})
        items = sorted({e["item"] for e in summary.get("rule_events", []) if e["equipment"] == name})
        ai = row.get("ai", {})
        if ai.get("flagged_records", 0) or ai.get("affected_hours", 0):
            items.append("AI 탐지")
        out[name] = {"status": "no_data" if not row.get("data_available") else "abnormal" if items else "normal",
                     "abnormal_items": items}
    return out


# -----------------------------------------------------------------------------
# 보고서 양식: 파싱(색 규칙) · 저장 · 복원 · 내보내기
# -----------------------------------------------------------------------------
def _fill(cell) -> str | None:
    try:
        if cell.fill is None or cell.fill.fill_type is None:
            return None
        rgb = cell.fill.fgColor.rgb
        return rgb.upper() if isinstance(rgb, str) and rgb not in ("00000000",) else None
    except AttributeError:
        return None


def classify_by_name(name: str) -> str:
    """색을 읽지 못한 항목의 분류. 애매하면 user."""
    n = _norm(name)
    if BLANK_NAME.search(name):
        return "blank"
    if re.search(r"설비명|설비$|대상설비|장비명|측정값|측정시각|판정|발생일시|이상내역|이상항목|기준값|편차|지속시간|전력|요금|탄소|부적합건수", n):
        return "auto"
    return "user"


def parse_xlsx(content: bytes) -> dict:
    """xlsx 양식 → {"fields": [...], "sheet": 이름}. 칸 배경색으로 auto/user/blank를 나눈다."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    top_left = {}
    for rng in ws.merged_cells.ranges:
        top_left[(rng.min_row, rng.min_col)] = True
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    top_left[(r, c)] = False
    fields: list[dict] = []
    section, header, table_field = None, None, None

    def is_value(cell):
        return _fill(cell) in AUTO_FILLS | USER_FILLS

    def text(cell):
        v = cell.value
        return str(v).strip() if v is not None and not str(v).startswith("=") and str(v).strip() else None

    for r in range(1, ws.max_row + 1):
        cells = [ws.cell(r, c) for c in range(1, ws.max_column + 1) if top_left.get((r, c), True)]
        labels = [c for c in cells if text(c) and not is_value(c) and _fill(c)]   # 칸 색이 있는 제목 칸
        values = [c for c in cells if is_value(c)]
        notes = [c for c in cells if text(c) and not _fill(c)]                     # 색 없는 안내문
        if labels and not values:
            if len(labels) == 1 and re.match(r"^\d+\.", text(labels[0]) or ""):
                section, header, table_field = re.sub(r"^\d+\.\s*", "", text(labels[0])), None, None
            elif len(labels) >= 3:
                header = {c.column: text(c) for c in labels}
                table_field = None
            continue
        if not values:
            continue
        first_value = min(c.column for c in values)
        if header:
            # 표 안에서는 머리글 첫 열(예: '점검 항목')에 있는 칸이 행 제목이다('210 V 이상' 같은 기준값 칸 제외)
            row_label = next((c for c in labels if c.column == min(header)), None)
        else:
            row_label = max((c for c in labels if c.column < first_value), key=lambda c: c.column, default=None)
        hint = next((text(c) for c in notes if c.column > first_value), None)
        if header and row_label is None and len(values) >= 3 and min(header) == first_value:
            # 행 제목이 없는 표(예: 이상 발생 내역) → 하나의 표 항목
            if table_field is None:
                table_field = {"name": section or "표", "source": "auto", "required": False, "kind": "table", "section": section,
                               "columns": [header[c] for c in sorted(header)], "rows": []}
                fields.append(table_field)
            table_field["rows"].append({header[c.column]: c.coordinate for c in values if c.column in header})
            continue
        if row_label is None:
            continue
        label = text(row_label)
        targets = values if header else values[:1]
        for cell in targets:
            formula = isinstance(cell.value, str) and cell.value.startswith("=")
            col_name = header.get(cell.column) if header and len(targets) > 1 else None
            name = f"{label}({col_name})" if col_name else label
            fill = _fill(cell)
            source = "auto" if fill in AUTO_FILLS else "user" if fill in USER_FILLS else classify_by_name(name)
            if BLANK_NAME.search(name):
                source = "blank"
            if formula:
                source = "formula"
            fields.append({"name": name, "source": source, "required": False, "cell": cell.coordinate, "section": section,
                           "kind": "formula" if formula else "field", "hint": hint if cell is targets[0] else None})
    return {"fields": fields, "sheet": ws.title}


def parse_docx(content: bytes) -> dict:
    """docx 양식 → 표의 '항목 | 값' 칸과 '항목:' 문단을 항목으로 추출(분류는 이름 기준)."""
    import docx
    document = docx.Document(io.BytesIO(content))
    names: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            for i in range(0, len(cells) - 1, 2):
                if cells[i] and cells[i] not in names and len(cells[i]) <= 30:
                    names.append(cells[i])
    for para in document.paragraphs:
        m = re.match(r"^\s*([^:：]{1,30})\s*[:：]", para.text)
        if m and m.group(1).strip() not in names:
            names.append(m.group(1).strip())
    return {"fields": [{"name": n, "source": classify_by_name(n), "required": False, "kind": "field"} for n in names]}


class TemplateStore:
    """브라우저 세션에만 양식과 업로드를 보관한다. 공용 파일은 읽기 전용이다."""

    def __init__(self, app_dir: Path = APP_DIR, state: dict | None = None):
        self.app_dir = Path(app_dir)
        self.state = state if state is not None else {}
        self.data = self.state.setdefault("_agent_templates", {})

    def ensure_default(self) -> dict:
        if "default" not in self.data:
            source = self.app_dir / DEFAULT_TEMPLATE_FILE
            content = source.read_bytes() if source.is_file() else b""
            parsed = parse_xlsx(content) if content else {"fields": [], "sheet": None}
            for f in parsed["fields"]:
                f["required"] = _norm(f["name"]) in DEFAULT_REQUIRED
                if _norm(f["name"]) == "종합판정":
                    f.update(source="user", required=True)
            self.data["default"] = {"name": DEFAULT_TEMPLATE_NAME, "is_default": True, "layout": "xlsx",
                "source_file": DEFAULT_TEMPLATE_FILE, "source_path": str(source),
                "source_sha256": hashlib.sha256(content).hexdigest(), "sheet": parsed["sheet"],
                "fields": parsed["fields"], "registered_at": datetime.now().isoformat(timespec="seconds")}
        self.data.setdefault("current", copy.deepcopy(self.data["default"]))
        return copy.deepcopy(self.data["default"])

    def get(self) -> dict:
        self.ensure_default()
        return copy.deepcopy(self.data["current"])

    def save(self, fields: list[dict], layout: str = "table", source_file: str | None = None,
             name: str | None = None, source_path: str | None = None, sheet: str | None = None) -> dict:
        clean = []
        for f in fields:
            fname = str(f.get("name", "")).strip()
            if not fname:
                continue
            source = f.get("source") if f.get("source") in ("auto", "user", "blank", "formula") else classify_by_name(fname)
            item = {k: v for k, v in f.items() if k in ("cell", "kind", "hint", "columns", "rows", "section") and v is not None}
            if _norm(fname) == "종합판정":
                source = "user"
            item.update(name=fname, source=source, required=bool(f.get("required", False)) or _norm(fname) == "종합판정", kind=f.get("kind") or "field")
            clean.append(item)
        if not clean:
            raise ValueError("양식 항목이 비어 있습니다.")
        template = {"name": name or (Path(source_file).stem if source_file else "사용자 양식"), "is_default": False,
            "layout": layout, "source_file": source_file, "source_path": None, "sheet": sheet,
            "fields": clean, "registered_at": datetime.now().isoformat(timespec="seconds")}
        if source_path:
            template["_source_bytes"] = self.data.get("uploads", {}).get(str(source_path))
        self.data["current"] = template
        self.mark_asked()
        return self.get()

    def reset(self) -> dict:
        self.data["current"] = self.ensure_default()
        self.mark_asked()
        return self.get()

    def store_upload(self, filename: str, content: bytes) -> Path:
        safe = re.sub(r"[^\w.\-가-힣]", "_", Path(filename).name)
        path = Path(hashlib.sha256(content).hexdigest()[:12] + "_" + safe)
        self.data.setdefault("uploads", {})[str(path)] = bytes(content)
        return path

    def template_asked(self) -> bool:
        return bool(self.data.get("template_asked"))

    def mark_asked(self) -> None:
        self.data["template_asked"] = True


def _safe_cell_text(value: Any) -> str:
    """사용자 입력이 수식으로 해석되지 않게 한다."""
    s = str(value)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def export_xlsx(draft: dict, template: dict) -> bytes:
    """원본 양식의 대표 내역과 누락 없는 전체 별첨을 함께 내보낸다."""
    import openpyxl
    if not draft.get("ready_to_export") or not draft.get("judgment_confirmed"):
        raise ValueError("필수 입력과 담당자 종합판정 확인을 완료한 뒤 다운로드하세요.")
    src = template.get("source_path")
    content = template.get("_source_bytes")
    if template.get("layout") == "xlsx" and (content or src and Path(src).is_file()):
        wb = openpyxl.load_workbook(io.BytesIO(content) if content else src)
        ws = wb[template["sheet"]] if template.get("sheet") in wb.sheetnames else wb.active
        for row in draft.get("rows", []):
            if row.get("source") in ("blank", "formula"):
                if row.get("source") == "formula" and "계약전력대비" in _norm(row["name"]) and row.get("cell"):
                    ws[row["cell"]] = None  # 최대수요전력이 없으므로 0%로 오해시키지 않음
                continue
            if row.get("kind") == "table":
                records = row.get("value") or []
                for i, cells in enumerate(row.get("rows") or []):
                    record = records[i] if i < len(records) else {}
                    for col, coord in cells.items():
                        value = record.get(col)
                        ws[coord] = _safe_cell_text(value) if value not in (None, "") else None
            elif row.get("cell"):
                value = row.get("value")
                ws[row["cell"]] = _safe_cell_text(value) if value not in (None, "") else None
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "보고서"
        ws.append([template.get("name", "보고서")])
        for row in draft.get("rows", []):
            if row.get("kind") != "table":
                ws.append([row["name"], "" if row.get("source") == "blank" else _safe_cell_text(row.get("value") or "")])
    evidence = wb.create_sheet("집계 기준 및 확인")
    evidence.append(["보고서 대상", f"{draft.get('date')} · {draft.get('equipment')}"])
    evidence.append(["담당자 확정 종합판정", draft.get("judgment")])
    evidence.append(["판정 확인 시각", draft.get("judgment_confirmed_at")])
    evidence.append(["집계 근거", draft.get("alert_basis")])
    evidence.append(["구간 의미", "설비·항목별 연속 탐지 구간이며 고장 사건 수가 아닙니다. 시간대 집계만 있으면 지속 시간은 확인 불가입니다."])
    evidence.append(["주의", "설정 기준과 AI 탐지는 별개입니다. 기준 이탈은 실제 고장이나 위험을 확정하지 않습니다."])
    for row in draft.get("rows", []):
        if row.get("kind") == "table":
            records = row.get("value") or []
            capacity = len(row.get("rows") or [])
            evidence.append([row["name"], f"본문 대표 {min(capacity, len(records))}개 / 전체 내역 {len(records)}개. 전체 내역 별첨을 확인하세요."])
    for title, events in (("설정 기준 전체 내역", draft.get("rule_events", [])), ("AI 탐지 전체 내역", draft.get("ai_events", []))):
        sheet = wb.create_sheet(title)
        sheet.append(["설비", "항목", "시작 또는 관측 시간대", "종료", "측정값", "단위", "설정 기준", "탐지 표본 수", "연속 구간 지속 시간", "집계 근거"])
        for e in events:
            sheet.append([_safe_cell_text(e.get("equipment", "")), _safe_cell_text(e.get("item", "")),
                str(e.get("timestamp", "")), str(e.get("end", "")), e.get("value"), e.get("unit"), e.get("threshold"),
                e.get("flagged_records"), duration_text(e), e.get("basis")])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 23
    evidence.column_dimensions["A"].width = 29
    evidence.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -----------------------------------------------------------------------------
# 도구 함수 12개
# -----------------------------------------------------------------------------
class AgentTools:
    """도구 함수 모음. state는 대시보드와 공유하는 st.session_state(또는 dict)."""

    def __init__(self, state: dict, templates: TemplateStore | None = None):
        self.state = state
        self.templates = templates or TemplateStore(state=state)
        state.setdefault("view_date", _iso(TODAY))
        state.setdefault("selected_equipment", EQUIPMENT_LIST[0][1])
        state.setdefault("thresholds", dict(DEFAULT_THRESHOLDS))
        state.setdefault("threshold_history", [])

    # -- 공용 ------------------------------------------------------------
    def _date(self, value: Any) -> date:
        d = _parse_date(value) if value else _parse_date(self.state.get("view_date"))
        if d is None:
            raise ValueError(f"날짜 형식을 확인하세요(YYYY-MM-DD): {value}")
        return d

    @property
    def thresholds(self) -> dict:
        th = self.state.get("thresholds") or DEFAULT_THRESHOLDS
        return {k: float(th.get(k, v)) for k, v in DEFAULT_THRESHOLDS.items()}

    def execute(self, name: str, args: dict | None) -> Any:
        """모델·데모 모드 공용 진입점. 실패해도 예외 대신 {"error": 사유}."""
        fn = getattr(self, name, None) if name in TOOL_NAMES else None
        if fn is None:
            return {"error": f"지원하지 않는 도구입니다: {name}"}
        try:
            return fn(**(args or {}))
        except TypeError as exc:
            return {"error": f"도구 인수를 확인하세요: {exc}"}
        except (ValueError, KeyError, OSError) as exc:
            return {"error": str(exc)}
        except Exception as exc:  # noqa: BLE001 — 도구 실패가 앱을 멈추지 않게
            return {"error": f"조회 중 오류가 발생했습니다: {type(exc).__name__}"}

    # 1
    def get_equipment_list(self) -> list[dict]:
        d = self._date(None)
        status = _status(d, self.thresholds)
        return [{"channel": ch, "name": name, "status": status[name]["status"],
                 "abnormal_items": status[name]["abnormal_items"], "date": _iso(d)} for ch, name in EQUIPMENT_LIST]

    # 2
    def get_anomalies(self, date: str | None = None, equipment: str | None = None) -> dict:
        d = self._date(date)
        eq = resolve_equipment(equipment)
        result = alert_summary(d, self.thresholds, eq)
        rule, ai = result.get("rule_events", []), result.get("ai_events", [])
        affected = []
        for row in result.get("equipment_rows", []):
            r, a = row.get("rule", {}), row.get("ai", {})
            if r.get("affected_hours", 0) or a.get("affected_hours", 0):
                affected.append({"equipment": row["equipment"], "rule_events": r.get("event_count"),
                    "ai_events": a.get("event_count"), "rule_affected_hours": r.get("affected_hours", 0),
                    "ai_affected_hours": a.get("affected_hours", 0),
                    "items": sorted({e["item"] for e in rule if e["equipment"] == row["equipment"]})})
        return {"date": _iso(d), "equipment": eq or "전체", "count": result.get("rule", {}).get("event_count"),
            "ai_count": result.get("ai", {}).get("event_count"), "rule": result.get("rule", {}), "ai": result.get("ai", {}),
            "items": rule[:20], "ai_items": ai[:10], "affected_equipment": affected,
            "data_available": result.get("data_available", False), "thresholds": self.thresholds,
            "basis": result.get("basis"), "note": result.get("note", ""),
            "shown_rule_items": min(len(rule), 20), "shown_ai_items": min(len(ai), 10),
            "event_count_note": "발생 구간은 설비·항목별 연속 표본이며 고장 사건 수가 아닙니다. 시간대만 있는 경우 구간 수·지속 시간은 확인 불가입니다."}

    # 3
    def get_measurements(self, date: str | None = None, equipment: str | None = None, item: str = "phase_voltage_R") -> dict:
        d = self._date(date)
        eq = resolve_equipment(equipment) or resolve_equipment(self.state.get("selected_equipment"))
        if item not in ITEMS:
            raise ValueError(f"item은 {', '.join(ITEMS)} 중 하나입니다.")
        col, label, unit, fam = ITEMS[item]
        th = self.thresholds[fam]
        g = _day(d, eq)
        if g.empty:
            return {"date": _iso(d), "equipment": eq, "item": item, "label": label, "unit": unit, "threshold": th,
                    "series": [], "data_available": False, "note": _no_data_note(d)}
        g = g.sort_values("hour")
        series = [{"hour": int(r.hour), "value": round(float(r[col]), 2)} for _, r in g.iterrows()]
        idx = g[col].idxmin()
        return {"date": _iso(d), "equipment": eq, "item": item, "label": label, "unit": unit, "threshold": th,
                "series": series, "min": round(float(g.loc[idx, col]), 2), "min_hour": int(g.loc[idx, "hour"]),
                "under_threshold_hours": [int(h) for h in g.loc[g[col] < th, "hour"]],
                "data_available": True, "basis": "시간별 값 = 그 시간 안의 5초 측정 최저값",
                "note": f"오늘은 기준 시점 {NOW:%H:%M}까지의 기록입니다." if d == NOW.date() and NOW.hour else ""}

    # 4
    def get_power(self, date: str | None = None, equipment: str | None = None, granularity: str = "hour") -> dict:
        d = self._date(date)
        eq = resolve_equipment(equipment)
        if granularity not in ("hour", "day"):
            raise ValueError("granularity는 hour 또는 day입니다.")
        g = _day(d, eq)
        kind = "actual"
        if g.empty:
            f = forecast()
            g = f[f["date"] == d] if len(f) else f
            if eq is not None and len(g):
                g = g[g["equipment"] == eq]
            kind = "forecast"
        if g.empty:
            return {"date": _iso(d), "equipment": eq or "전체", "series": [], "total_kwh": None, "total_cost_krw": None,
                    "total_co2_kg": None, "data_available": False, "note": _no_data_note(d)}
        per_hour = g.groupby("hour")["kwh"].sum()
        total = float(per_hour.sum())
        series = ([{"t": f"{h:02d}", **_money(float(v))} for h, v in per_hour.items()] if granularity == "hour"
                  else [{"t": _iso(d), **_money(total)}])
        money = _money(total)
        return {"date": _iso(d), "equipment": eq or "전체", "series": series,
                "total_kwh": money["kwh"], "total_cost_krw": money["cost_krw"], "total_co2_kg": money["co2_kg"],
                "peak_hour": f"{int(per_hour.idxmax()):02d}시", "peak_hour_kwh": round(float(per_hour.max()), 3),
                "hours": int(len(per_hour)), "cost_rate_krw_per_kwh": WON_PER_KWH, "co2_factor_kg_per_kwh": KGCO2_PER_KWH,
                "data_kind": kind, "data_available": True,
                "note": ("AI 예측값입니다(실측 아님). " if kind == "forecast" else "")
                        + (f"오늘은 기준 시점 {NOW:%H:%M}까지의 실측입니다." if d == NOW.date() and NOW.hour and kind == "actual" else "")}

    # 5
    def get_power_ranking(self, date: str | None = None) -> dict:
        d = self._date(date)
        g, kind = _day(d), "actual"
        if g.empty:
            f = forecast()
            g, kind = (f[f["date"] == d] if len(f) else f), "forecast"
        if g.empty:
            return {"date": _iso(d), "items": [], "total_kwh": None, "data_available": False, "note": _no_data_note(d)}
        per = g.groupby("equipment")["kwh"].sum().sort_values(ascending=False)
        total = float(per.sum())
        items = [{"rank": i + 1, "equipment": name, "kwh": round(float(v), 3),
                  "share_pct": round(float(v) / total * 100, 2) if total else None} for i, (name, v) in enumerate(per.items())]
        return {"date": _iso(d), "total_kwh": round(total, 3), "items": items, "data_kind": kind, "data_available": True,
                "note": ("AI 예측 기준 순위입니다. " if kind == "forecast" else "")
                        + (f"오늘은 {NOW:%H:%M}까지의 실측 기준입니다." if d == NOW.date() and NOW.hour and kind == "actual" else "")}

    # 6
    def get_forecast(self, month: str | None = None, equipment: str | None = None) -> dict:
        m = re.match(r"^(\d{4})-(\d{1,2})", str(month or self.state.get("view_date") or _iso(TODAY)))
        if not m:
            raise ValueError("month는 YYYY-MM 형식입니다.")
        year, mon = int(m.group(1)), int(m.group(2))
        eq = resolve_equipment(equipment)
        est = _month_estimate(year, mon, eq)
        if est is None:
            return {"month": f"{year}-{mon:02d}", "equipment": eq or "전체", "expected_kwh": None,
                    "data_available": False, "note": "해당 월의 실측·예측 자료가 없습니다."}
        prev = pd.Timestamp(year, mon, 1) - pd.offsets.MonthBegin(1)
        prev_est = _month_estimate(prev.year, prev.month, eq)
        diff = round((est["total"] - prev_est["total"]) / prev_est["total"] * 100, 1) if prev_est and prev_est["total"] else None
        kind = "actual" if est["pred_hours"] == 0 and est["filled_hours"] == 0 else (
            "forecast" if est["actual_hours"] == 0 else "actual+forecast")
        meta = _model_meta()
        money = _money(est["total"])
        return {"month": f"{year}-{mon:02d}", "equipment": eq or "전체",
                "expected_kwh": money["kwh"], "expected_cost_krw": money["cost_krw"], "expected_co2_kg": money["co2_kg"],
                "actual_kwh_so_far": round(est["actual"], 3), "actual_hours": est["actual_hours"],
                "forecast_kwh": round(est["pred"], 3), "forecast_hours": est["pred_hours"],
                "filled_hours_by_average": est["filled_hours"],
                "prev_month_diff_pct": diff, "mae_kwh_per_hour": meta["mae_kwh_per_hour"] if eq is None else None,
                "model": meta["model"], "data_kind": kind, "data_available": True,
                "cost_rate_krw_per_kwh": WON_PER_KWH, "co2_factor_kg_per_kwh": KGCO2_PER_KWH,
                "note": "월 전체 추정 = 시간별 실측(기준 시점까지) + 이후 AI 예측 + 자료가 없는 시간은 평균으로 보정. "
                        "MAE는 공장 전체 기준이며 설비별 오차는 제공하지 않습니다."}

    # 7
    def get_thresholds(self) -> dict:
        th = self.thresholds
        return {**th, "is_default": th == DEFAULT_THRESHOLDS, "changed_log": self.state.get("threshold_history", [])[-10:],
                "note": "규칙 임계치입니다. AI 이상탐지 기준은 모델에 고정되어 바뀌지 않습니다."}

    # 8
    def set_thresholds(self, phase_voltage: float | None = None, line_voltage: float | None = None,
                       power_factor: float | None = None, reason: str = "") -> dict:
        before = self.thresholds
        after = dict(before)
        limits = {"phase_voltage": (100, 300), "line_voltage": (200, 500), "power_factor": (1, 100)}
        for key, value in (("phase_voltage", phase_voltage), ("line_voltage", line_voltage), ("power_factor", power_factor)):
            if value is None:
                continue
            v = float(value)
            lo, hi = limits[key]
            if not lo <= v <= hi:
                raise ValueError(f"{FAMILY_LABEL[key]} 임계치는 {lo}~{hi} 범위로 입력하세요.")
            after[key] = v
        if after == before:
            return {"before": before, "after": after, "changed": False, "note": "바뀐 값이 없습니다."}
        d = self._date(None)
        def abnormal(th):
            return {n for n, s in _status(d, th).items() if s["status"] == "abnormal"}
        available = not _day(d).empty
        b_summary = alert_summary(d, before)
        a_summary = alert_summary(d, after)
        was, now = (abnormal(before), abnormal(after)) if available else (set(), set())
        change = ", ".join(f"{FAMILY_LABEL[k]} {before[k]:g} → {after[k]:g}" for k in after if after[k] != before[k])
        entry = {"at": datetime.now().strftime("%Y-%m-%d %H:%M"), "by": "사용자", "change": change,
                 "reason": str(reason or "")[:200], "before": before, "after": after}
        self.state["thresholds"] = after
        self.state.setdefault("threshold_history", []).append(entry)
        return {"before": before, "after": after, "changed": True, "analysis_date": _iso(d), "data_available": available,
                "anomaly_count_before": b_summary.get("rule", {}).get("event_count") if available else None,
                "anomaly_count_after": a_summary.get("rule", {}).get("event_count") if available else None,
                "newly_abnormal": sorted(now - was) if available else None,
                "no_longer_abnormal": sorted(was - now) if available else None,
                "log_entry": entry, "note": "규칙 임계치만 바꿨습니다. AI 판정은 바뀌지 않습니다."}

    # 9
    def get_report_template(self) -> dict:
        t = self.templates.get()
        return {"name": t.get("name"), "is_default": t.get("is_default", False), "layout": t.get("layout"),
                "source_file": t.get("source_file"), "registered_at": t.get("registered_at"),
                "template_asked": self.templates.template_asked(),
                "fields": [{"name": f["name"], "source": f["source"], "required": f.get("required", False)}
                           for f in t.get("fields", [])]}

    # 10
    def save_report_template(self, fields: list, layout: str = "table", source_file: str | None = None) -> dict:
        pending = self.state.get("pending_template") or {}
        # 업로드 파일을 확인한 경우에는 원본 칸 위치(cell)를 함께 저장한다
        by_name = {_norm(f.get("name")): f for f in pending.get("fields", [])}
        merged = []
        for f in fields:
            f = {"name": f} if isinstance(f, str) else dict(f)
            merged.append({**by_name.get(_norm(f.get("name")), {}), **f})
        t = self.templates.save(merged, layout=pending.get("layout", layout), source_file=source_file or pending.get("source_file"),
                                source_path=pending.get("source_path"), sheet=pending.get("sheet"))
        self.state["pending_template"] = None
        self._refresh_draft()
        return {"saved": True, "name": t["name"], "is_default": False, "field_count": len(t["fields"]),
                "fields": [{"name": f["name"], "source": f["source"], "required": f["required"]} for f in t["fields"]]}

    # 11
    def reset_report_template(self) -> dict:
        t = self.templates.reset()
        self.state["pending_template"] = None
        self._refresh_draft()
        return {"reset": True, "name": t["name"], "is_default": True, "field_count": len(t["fields"])}

    # 12
    def create_report(self, date: str | None = None, equipment: str | None = None,
                      action_taken: str = "", extra_fields: Any = None) -> dict:
        draft = self.state.get("report_draft") or {}
        d = self._date(date or draft.get("date"))
        eq = resolve_equipment(equipment) or draft.get("equipment") or resolve_equipment(self.state.get("selected_equipment"))
        same = bool(draft and draft.get("date") == _iso(d) and draft.get("equipment") == eq)
        if draft and not same:
            return {"error": "작성 중인 보고서의 대상이 다릅니다. 기존 초안을 유지했습니다. 화면의 '현재 선택으로 새 보고서'에서 확인 후 시작하세요.",
                    "target_conflict": True, "draft_date": draft.get("date"), "draft_equipment": draft.get("equipment")}
        user_values = dict(draft.get("user_values", {})) if same else {}
        template = self.templates.get()
        names = {_norm(f["name"]): f["name"] for f in template.get("fields", []) if f["source"] == "user"}
        extras = extra_fields or {}
        if isinstance(extras, list):
            extras = {str(r.get("name")): r.get("value") for r in extras if isinstance(r, dict)}
        unknown = [str(k) for k in extras if _norm(k) not in names] if isinstance(extras, dict) else []
        if unknown:
            return {"error": "양식에서 찾지 못한 항목입니다. 항목 이름을 확인하세요.", "unknown_fields": unknown,
                    "allowed_fields": list(names.values())}
        for key, value in (extras.items() if isinstance(extras, dict) else []):
            target = names[_norm(key)]
            if value is None or not str(value).strip():
                user_values.pop(target, None)
            else:
                user_values[target] = str(value)
        if action_taken:
            action_field = next((n for k, n in names.items() if "조치내용" in k), None)
            if action_field:
                user_values[action_field] = str(action_taken)
        action_field = next((n for k, n in names.items() if "조치내용" in k), None)
        action = user_values.get(action_field, "") if action_field else (action_taken or draft.get("action_taken", ""))
        new_draft = self._build_draft(d, eq, user_values, template, action)
        new_draft["revision"] = int(draft.get("revision", 0)) + 1
        self.state["report_draft"] = new_draft
        self._set_export_state()
        result = copy.deepcopy(new_draft)
        result.pop("rule_events", None)
        result.pop("ai_events", None)
        result["unknown_fields"] = []
        result["template_asked"] = self.templates.template_asked()
        for row in result["rows"]:
            row.pop("rows", None)
            row.pop("cell", None)
            if isinstance(row.get("value"), list):
                row["total_count"] = len(row["value"])
                row["value"] = row["value"][:3]
                row["shown_count"] = len(row["value"])
        return result

    def start_new_report(self, date: str | None = None, equipment: str | None = None) -> dict:
        """사용자가 화면에서 새 초안을 선택했을 때만 호출한다. 모델 도구에는 노출하지 않는다."""
        d = self._date(date)
        eq = resolve_equipment(equipment) or resolve_equipment(self.state.get("selected_equipment"))
        previous = self.state.get("report_draft")
        if previous:
            self.state["previous_report_draft"] = copy.deepcopy(previous)
        self.state.pop("report_draft", None)
        return self.create_report(_iso(d), eq)

    def confirm_judgment(self, judgment: str) -> dict:
        """화면의 담당자 확인 버튼 전용. 모델 호출 가능한 도구 목록에 포함하지 않는다."""
        if judgment not in ("적합", "추가 점검 필요", "조치 후 재확인 필요"):
            raise ValueError("종합판정을 선택하세요.")
        draft = self.state.get("report_draft")
        if not draft:
            raise ValueError("보고서 초안을 먼저 시작하세요.")
        draft["judgment"] = judgment
        draft["judgment_confirmed"] = True
        draft["judgment_confirmed_at"] = datetime.now().isoformat(timespec="seconds")
        for row in draft["rows"]:
            if _norm(row["name"]) == "종합판정":
                row["value"] = judgment
                draft["user_values"][row["name"]] = judgment
        self._set_export_state()
        return {"confirmed": True, "ready_to_export": draft["ready_to_export"]}

    def _set_export_state(self) -> None:
        draft = self.state.get("report_draft")
        if not draft:
            return
        missing = [r["name"] for r in draft["rows"] if r.get("required") and r.get("source") == "user" and not str(r.get("value") or "").strip()]
        blockers = missing + list(draft.get("required_auto_missing", []))
        if not draft.get("judgment_confirmed"):
            blockers.append("담당자 종합판정 확인")
        if not draft.get("data_available"):
            blockers.append("대상 날짜의 실측 자료 확인")
        draft["missing_user_fields"] = missing
        draft["export_blockers"] = list(dict.fromkeys(blockers))
        draft["ready_to_export"] = not blockers

    # -- 보고서 초안 ------------------------------------------------------
    def _refresh_draft(self) -> None:
        draft = self.state.get("report_draft")
        if draft:
            d = _parse_date(draft["date"])
            self.state["report_draft"] = self._build_draft(d, draft["equipment"], draft.get("user_values", {}),
                                                           self.templates.get(), draft.get("action_taken", ""))
            self.state["report_draft"]["revision"] = int(draft.get("revision", 0)) + 1
            self._set_export_state()

    def _context(self, d: date, eq: str | None) -> dict:
        g = _day(d, eq)
        th = self.thresholds
        meas = {}
        for key, (col, label, unit, fam) in ITEMS.items():
            if g.empty:
                continue
            idx = g[col].idxmin()
            value = float(g.loc[idx, col])
            meas[key] = {"value": f"{value:.1f} {unit}", "time": f"{d} {int(g.loc[idx, 'hour']):02d}시",
                         "verdict": "적합" if value >= th[fam] else "부적합"}
        def power(day):
            p = self.get_power(_iso(day), eq, "day")
            return p if p.get("data_available") and p.get("data_kind") == "actual" else None
        prev_day = (pd.Timestamp(d) - pd.DateOffset(months=1)).date()
        summary = alert_summary(d, th, eq)
        rule, ai = summary.get("rule_events", []), summary.get("ai_events", [])
        return {"date": d, "equipment": eq, "meas": meas, "power": power(d), "prev": power(prev_day),
                "events": rule, "ai_events": ai, "alert_summary": summary}

    def _auto_value(self, field: dict, ctx: dict) -> Any:
        name = _norm(field["name"])
        base, _, suffix = name.partition("(")
        suffix = suffix.rstrip(")")
        if field.get("kind") == "table":
            out = []
            for e in ctx["events"]:
                out.append({"발생 일시": e["timestamp"], "이상 항목": f"{e['item']}", "측정값": f"{e['value']} {e['unit']}",
                            "기준값": f"{e['threshold']:g} {e['unit']}", "편차": f"{e.get('deviation_pct', round((e['value'] - e['threshold']) / e['threshold'] * 100, 1))}%",
                            "지속 시간": duration_text(e)})
            return out or None
        if re.fullmatch(r"(설비명|설비|대상설비|점검대상설비|장비명)", base):
            return ctx["equipment"]
        item_map = {"r상전압": "phase_voltage_R", "s상전압": "phase_voltage_S", "t상전압": "phase_voltage_T",
                    "rs선간전압": "line_voltage_RS", "st선간전압": "line_voltage_ST", "tr선간전압": "line_voltage_TR",
                    "r역률": "power_factor_R", "s역률": "power_factor_S", "t역률": "power_factor_T"}
        if base in item_map:
            m = ctx["meas"].get(item_map[base])
            if not m:
                return None
            return m["time"] if "시각" in suffix else m["verdict"] if "판정" in suffix else m["value"]
        if base == "부적합건수":
            return f"{sum(1 for m in ctx['meas'].values() if m['verdict'] == '부적합')} 건" if ctx["meas"] else None
        p = ctx["prev"] if "전월" in suffix else ctx["power"]
        if re.search(r"전력소비량|전력량|전력사용량", base):
            return f"{p['total_kwh']:,.1f}" if p else None
        if re.search(r"전기요금|요금", base):
            return f"{p['total_cost_krw']:,.0f}" if p else None
        if re.search(r"탄소", base):
            return f"{p['total_co2_kg']:,.1f}" if p else None
        if re.search(r"최대수요전력|계약전력", base):
            return None  # 계약전력·최대수요전력의 측정 근거가 없음
        first = ctx["events"][0] if ctx["events"] else None
        if re.search(r"발생일시", base):
            return first["timestamp"] if first else None
        if re.search(r"이상내역|이상항목", base):
            return f"{first['item']} {first['value']}{first['unit']} (기준 {first['threshold']:g}{first['unit']})" if first else None
        if re.search(r"측정값", base):
            return f"{first['value']} {first['unit']}" if first else None
        if re.search(r"기준값|임계", base):
            return f"{first['threshold']:g} {first['unit']}" if first else None
        if re.search(r"편차", base):
            return f"{first.get('deviation_pct', round((first['value'] - first['threshold']) / first['threshold'] * 100, 1))}%" if first else None
        if re.search(r"지속시간", base):
            return duration_text(first) if first else None
        if re.search(r"점검일|일자|날짜", base):
            return _iso(ctx["date"])
        return None

    def _build_draft(self, d: date, eq: str | None, user_values: dict, template: dict, action_taken: str) -> dict:
        ctx = self._context(d, eq)
        rows, missing_user, missing_auto = [], [], []
        for f in template.get("fields", []):
            row = {k: f.get(k) for k in ("name", "source", "required", "kind", "hint", "cell", "rows", "columns", "section")}
            if f["source"] == "user":
                row["value"] = user_values.get(f["name"])
                if f.get("required") and not row["value"]:
                    missing_user.append(f["name"])
            elif f["source"] in ("auto", "formula"):
                row["value"] = self._auto_value(f, ctx)
                if f["source"] == "formula" and row["value"] is None:
                    row["value"] = "(엑셀에서 자동 계산)"
                if f["source"] == "auto" and row["value"] in (None, "", []):
                    missing_auto.append(f["name"])
            else:
                row["value"] = None
            rows.append(row)
        required_auto_missing = [f["name"] for f in template.get("fields", [])
                                 if f["source"] == "auto" and f.get("required") and f["name"] in missing_auto]
        return {"date": _iso(d), "equipment": eq, "template_name": template.get("name"),
                "is_default_template": template.get("is_default", False), "rows": rows, "user_values": user_values,
                "action_taken": action_taken, "missing_user_fields": missing_user, "missing_auto_fields": missing_auto,
                "ready_to_export": False, "required_auto_missing": required_auto_missing,
                "export_blockers": missing_user + required_auto_missing + ["담당자 종합판정 확인"],
                "rule_events": ctx["events"], "ai_events": ctx["ai_events"], "alert_basis": ctx["alert_summary"].get("basis"),
                "judgment_confirmed": False, "judgment": None,
                "data_available": not _day(d, eq).empty,
                "note": "" if not _day(d, eq).empty else _no_data_note(d)}


TOOL_NAMES = ("get_equipment_list", "get_anomalies", "get_measurements", "get_power", "get_power_ranking",
              "get_forecast", "get_thresholds", "set_thresholds", "get_report_template", "save_report_template",
              "reset_report_template", "create_report")


def _fn(name: str, description: str, properties: dict, required: list[str] | None = None) -> dict:
    return {"type": "function", "function": {"name": name, "description": description, "parameters": {
        "type": "object", "properties": properties, "required": required or [], "additionalProperties": False}}}


DATE = {"type": "string", "description": "YYYY-MM-DD. 생략하면 대시보드에서 선택한 조회 날짜"}
EQUIP = {"type": ["string", "null"], "description": "설비 이름 예: '12 (4호기)', '4호기'. null이면 공장 전체"}

TOOL_SCHEMAS = [
    _fn("get_equipment_list", "설비 13개의 채널 번호·이름·조회 날짜 상태(정상/이상/자료없음)와 이상 항목.", {}),
    _fn("get_anomalies", "그 날의 이상 목록(규칙 임계치 미달 구간 + AI 이상탐지). 실측일만 가능.",
        {"date": DATE, "equipment": EQUIP}),
    _fn("get_measurements", "한 설비·한 항목의 시간별 값(그 시간 최저값)과 임계치, 최저값, 임계치를 밑돈 시각.",
        {"date": DATE, "equipment": EQUIP, "item": {"type": "string", "enum": list(ITEMS)}}, ["item"]),
    _fn("get_power", "전력 사용량(kWh)·전기요금(원)·탄소(kgCO2). 실측일은 실측, 기준 시점 이후 날짜는 AI 예측(data_kind로 구분).",
        {"date": DATE, "equipment": EQUIP, "granularity": {"type": "string", "enum": ["hour", "day"]}}),
    _fn("get_power_ranking", "그 날 설비별 전력 소비 순위 13개와 비중(%).", {"date": DATE}),
    _fn("get_forecast", "그 달의 예상 전력량·요금·탄소(실측+AI 예측), 전월 대비 증감률, 예측 오차(MAE).",
        {"month": {"type": "string", "description": "YYYY-MM. 생략하면 조회 날짜의 달"}, "equipment": EQUIP}),
    _fn("get_thresholds", "현재 규칙 임계치 세 개와 기본값 여부, 변경 이력.", {}),
    _fn("set_thresholds", "사용자가 명시적으로 요청했을 때만 규칙 임계치를 바꾸고 조회 날짜로 재분석한다.",
        {"phase_voltage": {"type": ["number", "null"]}, "line_voltage": {"type": ["number", "null"]},
         "power_factor": {"type": ["number", "null"]}, "reason": {"type": "string"}}),
    _fn("get_report_template", "지금 쓰는 보고서 양식(항목·auto/user/blank 분류)과 최초 양식 질문 여부.", {}),
    _fn("save_report_template", "사용자가 확인한 양식 항목을 저장한다. 확인받기 전에는 호출하지 않는다.",
        {"fields": {"type": "array", "items": {"type": "object", "properties": {
            "name": {"type": "string"}, "source": {"type": "string", "enum": ["auto", "user", "blank"]},
            "required": {"type": "boolean"}}, "required": ["name", "source"]}},
         "layout": {"type": "string"}, "source_file": {"type": ["string", "null"]}}, ["fields"]),
    _fn("reset_report_template", "기본 양식(제조현장 안전 점검 보고서)으로 되돌린다.", {}),
    _fn("create_report", "등록된 양식으로 보고서 초안을 만들거나 갱신한다. 사용자가 말한 조치 내용은 action_taken에 원문 그대로, "
                         "사용자가 알려준 항목 값은 extra_fields({항목 이름: 값})에 넣는다.",
        {"date": DATE, "equipment": EQUIP, "action_taken": {"type": "string"},
         "extra_fields": {"type": ["object", "null"], "additionalProperties": {"type": "string"}}}),
]
