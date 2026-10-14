"""Integrated report service: root asset, gates, source identity and preserved form."""
import hashlib
import io
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))
from agent_runtime.agent_reports import ReportService  # noqa: E402


@pytest.fixture
def service():
    base = APP / "runtime" / "test_reports"
    base.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=base) as temp:
        yield ReportService(Path(temp))


def complete(service, **kwargs):
    extra = {"점검 일시": "2025-04-29 10:30", "점검자": "검증 담당", "종합 판정": "주의"}
    extra.update(kwargs.pop("extra_fields", {}))
    return service.create_report("2025-04-29", "13(3호기)", extra_fields=extra, **kwargs)


def test_root_asset_is_original_and_export_preserves_form(service):
    assert service.asset == APP / "BasicSafetyReportForm.xlsx"
    content = service.asset.read_bytes()
    git_blob = hashlib.sha1(f"blob {len(content)}\0".encode() + content).hexdigest()
    assert git_blob == "a748e7a59a633420a5358410e65aa957180c9c27"
    output = openpyxl.load_workbook(io.BytesIO(service.export_xlsx(complete(service))))
    sheet = output["안전점검보고서"]
    assert len(sheet.merged_cells.ranges) == 73
    assert all(sheet[cell].data_type == "f" for cell in ("C26", "C35", "E31", "E32", "E33", "E34"))
    assert all(sheet[cell].value is None for cell in ("G1", "G2", "H2", "I2", "D61", "F61", "H61", "D62", "F62", "H62"))
    assert service.asset.read_bytes() == content


def test_required_missing_data_is_listed_and_not_fabricated(service):
    service.save_report_template([{"name": "R상 전압 측정값", "source": "auto", "required": True, "maps_to": "measurements.phase_voltage_R.value"}])
    result = complete(service, measurements={"phase_voltage_R": {"data_available": False}})
    assert result["missing_auto_fields"] == ["R상 전압 측정값"]
    assert "R상 전압 측정값" in result["missing_user_fields"]
    assert not result["ready_to_export"]
    with pytest.raises(ValueError, match="R상 전압"):
        service.export_xlsx(result)


def test_optional_absent_data_does_not_create_values_or_block_human_report(service):
    result = complete(service, power={"total_kwh": None, "data_available": False})
    values = {r["name"]: r["value"] for r in result["rows"]}
    assert values["당일 전력 소비량 (kWh)"] is None
    assert "당일 전력 소비량 (kWh)" in result["missing_auto_fields"]
    assert result["ready_to_export"]


def test_action_gate_and_formula_injection_protection(service):
    result = complete(service, action_taken="콘덴서 교체", extra_fields={"점검자": "=1+1"})
    assert set(result["missing_user_fields"]) == {"조치 일시", "조치자", "조치 후 확인"}
    with pytest.raises(ValueError):
        service.export_xlsx(result)
    good = complete(service, action_taken="콘덴서 교체", extra_fields={"점검자": "=1+1", "조치 일시": "2025-04-29 11:00", "조치자": "김담당", "조치 후 확인": "현장 재점검"})
    output = openpyxl.load_workbook(io.BytesIO(service.export_xlsx(good)))
    assert output["안전점검보고서"]["C7"].value == "'=1+1"


def test_sources_exact_observation_and_all_events_survive_export(service):
    events = [{"timestamp": f"2025-04-29T0{i}:00:00", "item": "voltageR", "value": 199, "threshold": 210, "duration_seconds": 15} for i in range(6)]
    result = complete(service, measurements={"phase_voltage_R": {"representative_value": 199, "representative_timestamp": "2025-04-29T00:12:05", "threshold": 210, "source_id": "same-id"}}, power={"total_kwh": 10, "source_id": "same-id", "peak_kw": 2}, anomalies={"items": events, "source_id": "same-id", "ai_items": [{**events[0], "item": "AI"}]})
    output = openpyxl.load_workbook(io.BytesIO(service.export_xlsx(result)))
    sheet = output["안전점검보고서"]
    assert sheet["E17"].value == "2025-04-29T00:12:05"
    assert sheet["C32"].value == 1800
    assert sheet["C34"].value is None
    assert output["이상 내역 별지"].max_row == 8
    assert "AI 탐지 구간" in output.sheetnames


def test_mixed_source_report_cannot_be_exported(service):
    result = complete(service, power={"data_id": "A"}, anomalies={"source_id": "B", "items": []})
    assert not result["ready_to_export"]
    with pytest.raises(ValueError, match="식별자"):
        service.export_xlsx(result)


def test_template_preview_needs_confirmation_and_reset_restores_default(service):
    before = service.template_path.read_bytes()
    preview = service.parse_template("BasicSafetyReportForm.xlsx", service.asset.read_bytes())
    assert preview["preview_only"]
    assert service.template_path.read_bytes() == before
    saved = service.save_report_template([{"name": "현장 메모", "source": "user", "required": False}])
    assert not saved["is_default"] and saved["layout"] == "table"
    assert service.reset_report_template()["is_default"]
