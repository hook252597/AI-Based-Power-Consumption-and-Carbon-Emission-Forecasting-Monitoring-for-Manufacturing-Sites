"""Local report drafts and XLSX downloads; templates never execute uploaded content.

Only template preferences are persisted. Report contents stay in the caller's session.
The supplied default form is retained; custom forms export as an explicit field table.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties


DEFAULT_DIR = Path(__file__).resolve().parent
TEAL, YELLOW = "FFEAF3F2", "FFFFF9E0"
CORE = ("점검 일시", "설비명", "점검자", "종합 판정")
ACTION = ("조치 내용", "조치 일시", "조치자", "조치 후 확인")
ALIASES = {
    "점검 일시": ("inspection_at", "inspection_datetime", "점검일시"),
    "점검자": ("inspector",), "설비명": ("equipment", "설비"),
    "종합 판정": ("overall_judgement", "overall_judgment", "종합판정"),
    "조치 내용": ("action_taken", "조치내용"), "조치 일시": ("action_at", "조치일시"),
    "조치자": ("action_by",), "조치 후 확인": ("action_confirmation", "조치후확인"),
    "계약전력 (kW)": ("contract_power_kw", "계약전력"),
}
METRICS = (
    ("phase_voltage_R", "R상 전압", "voltageR", 210, "V"),
    ("phase_voltage_S", "S상 전압", "voltageS", 210, "V"),
    ("phase_voltage_T", "T상 전압", "voltageT", 210, "V"),
    ("line_voltage_RS", "RS 선간전압", "voltageRS", 360, "V"),
    ("line_voltage_ST", "ST 선간전압", "voltageST", 360, "V"),
    ("line_voltage_TR", "TR 선간전압", "voltageTR", 360, "V"),
    ("power_factor_R", "R 역률", "powerFactorR", 85, "%"),
    ("power_factor_S", "S 역률", "powerFactorS", 85, "%"),
    ("power_factor_T", "T 역률", "powerFactorT", 85, "%"),
)


def _key(value: Any) -> str:
    return re.sub(r"[\s_()·:/-]+", "", str(value)).casefold()


def _canonical(name: str) -> str:
    for canonical, aliases in ALIASES.items():
        if _key(name) in {_key(canonical), *map(_key, aliases)}:
            return canonical
    return name.strip()


def _blank_field(name: str) -> bool:
    return any(word in _key(name) for word in ("문서번호", "결재", "서명", "승인자", "검토자", "작성자"))


def _missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _safe_text(value: Any) -> Any:
    """Force user strings to literal cells, never formulas or external links."""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _field(name: str, source: str, cell: str | None = None, **kwargs: Any) -> dict:
    value = {"name": name, "source": source, "required": name in CORE}
    if cell:
        value["cell"] = cell
    value.update(kwargs)
    return value


def _default_fields() -> list[dict]:
    fields = [_field("문서번호", "blank", "G1"), _field("결재란", "blank", "G2")]
    for name, cell in (
        ("점검 일시", "C5"), ("점검 구분", "C6"), ("점검자", "C7"),
        ("설비 위치", "C11"), ("정격 용량 (kW)", "C12"), ("정격 전압 / 전류", "C13"),
        ("계약전력 (kW)", "C29"), ("조치 내용", "C45"), ("조치 일시", "C48"),
        ("조치자", "C49"), ("조치 후 확인", "C50"), ("종합 판정", "C53"),
        ("종합 의견", "C54"), ("차기 점검 예정일", "C57"),
    ):
        fields.append(_field(name, "user", cell))
    fields.append(_field("설비명", "auto", "C10", maps_to="equipment"))
    for row, (item, label, _raw, _threshold, _unit) in enumerate(METRICS, 17):
        for col, suffix, attr in (("D", "측정값", "value"), ("E", "측정 시각", "timestamp"), ("F", "판정", "judgement")):
            fields.append(_field(f"{label} {suffix}", "auto", f"{col}{row}", maps_to=f"measurements.{item}.{attr}"))
        fields.append(_field(f"{label} 비고", "user", f"G{row}"))
    fields.append(_field("부적합 건수", "auto", "C26", maps_to="nonconforming_count", formula=True))
    for row, (label, attr) in enumerate((("전력 소비량 (kWh)", "total_kwh"), ("전기 요금 (원)", "total_cost_krw"), ("탄소 배출량 (kgCO2)", "total_co2_kg"), ("최대 수요전력 (kW)", "max_demand_kw")), 31):
        fields.extend([
            _field(f"당일 {label}", "auto", f"C{row}", maps_to=f"power.{attr}"),
            _field(f"전월 동일 {label}", "auto", f"D{row}", maps_to=f"previous_power.{attr}"),
            _field(f"{label} 증감", "auto", f"E{row}", maps_to=f"changes.{attr}", formula=True),
            _field(f"{label} 비고", "user", f"F{row}"),
        ])
    fields.append(_field("계약전력 대비", "auto", "C35", maps_to="contract_ratio", formula=True))
    fields.append(_field("이상 발생 내역", "auto", maps_to="anomaly_summary"))
    for i in range(4):
        fields.append(_field(f"이상 내역 {i + 1} 비고", "user", f"H{39 + i}"))
    for name, col in (("작성자", "D"), ("검토자", "F"), ("승인자", "H")):
        fields.extend([_field(f"{name} 성명", "blank", f"{col}61"), _field(f"{name} 서명 / 일자", "blank", f"{col}62")])
    return fields


def _lookup(data: dict, path: str) -> Any:
    result: Any = data
    for key in path.split("."):
        if not isinstance(result, dict):
            return None
        result = result.get(key)
    return result


class ReportService:
    def __init__(self, app_dir: Path = DEFAULT_DIR):
        self.app_dir = Path(app_dir)
        self.runtime = self.app_dir / "runtime"
        self.runtime.mkdir(parents=True, exist_ok=True)
        self.asset = self.app_dir / "BasicSafetyReportForm.xlsx"
        if not self.asset.exists():
            self.asset = DEFAULT_DIR / "BasicSafetyReportForm.xlsx"
        self.template_path = self.runtime / "report_template.json"
        self.default_path = self.runtime / "report_template_default.json"
        self.default = {
            "name": "제조현장 안전 점검 보고서", "is_default": True, "layout": "original_xlsx",
            "source_file": "BasicSafetyReportForm.xlsx", "fields": _default_fields(),
            "note": "기본 Excel 양식 유지. 문서번호·결재·서명은 색보다 공란 규칙을 우선합니다.",
        }
        if not self.default_path.exists():
            self._write(self.default_path, self.default)
        if not self.template_path.exists():
            self._write(self.template_path, self.default)

    @staticmethod
    def _write(path: Path, value: dict) -> None:
        temp = path.with_suffix(".tmp")
        temp.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        temp.replace(path)

    def get_report_template(self) -> dict:
        try:
            value = json.loads(self.template_path.read_text(encoding="utf-8"))
            if not isinstance(value, dict) or not isinstance(value.get("fields"), list):
                raise ValueError("fields missing")
            return value
        except (OSError, ValueError, TypeError):
            return copy.deepcopy(self.default)

    def save_report_template(self, fields: list[dict], layout: str = "table", source_file: str | None = None) -> dict:
        if not fields or len(fields) > 200:
            raise ValueError("양식 항목은 1~200개로 지정해 주세요.")
        normalized, seen = [], set()
        default_by_name = {f["name"]: f for f in self.default["fields"]}
        for raw in fields:
            name = _canonical(str(raw.get("name", "")))
            if not name or len(name) > 120:
                raise ValueError("항목 이름은 1~120자로 지정해 주세요.")
            if _key(name) in seen:
                raise ValueError(f"중복 항목: {name}")
            seen.add(_key(name))
            source = raw.get("source", "user")
            if source not in {"auto", "user", "blank"}:
                raise ValueError(f"지원하지 않는 입력 방식: {source}")
            if _blank_field(name):
                source = "blank"
            known = default_by_name.get(name, {})
            mapping = raw.get("maps_to") or known.get("maps_to")
            if source == "auto" and not mapping:
                source = "user"
            field = _field(name, source, required=bool(raw.get("required", name in CORE)) and source != "blank")
            if name in CORE:
                field["required"] = True
                field["source"] = "auto" if name == "설비명" else "user"
                if name == "설비명":
                    mapping = "equipment"
            if mapping and field["source"] == "auto":
                field["maps_to"] = str(mapping)
            normalized.append(field)
        # The agreed four checks apply even when an uploaded form omits a field.
        for name in CORE:
            if _key(name) not in seen:
                normalized.append(_field(name, "auto" if name == "설비명" else "user", maps_to="equipment" if name == "설비명" else ""))
        value = {
            "name": "사용자 보고서 양식", "is_default": False, "layout": "table", "fields": normalized,
            "source_file": Path(source_file).name if source_file else None,
            "registered_at": datetime.now().isoformat(timespec="seconds"),
            "note": "확정한 항목을 간단한 Excel 표로 내보냅니다. 업로드 파일의 원래 디자인은 복제하지 않습니다.",
        }
        self._write(self.template_path, value)
        return copy.deepcopy(value)

    def reset_report_template(self) -> dict:
        self._write(self.template_path, self.default)
        return copy.deepcopy(self.default)

    def parse_template(self, filename: str, content: bytes) -> dict:
        """Return a preview only; caller must explicitly confirm before saving."""
        suffix = Path(filename).suffix.lower()
        if suffix not in {".xlsx", ".docx"}:
            raise ValueError(".xlsx 또는 .docx 양식을 선택하거나 항목을 직접 입력해 주세요.")
        if not content or len(content) > 10 * 1024 * 1024:
            raise ValueError("양식 파일은 10 MB 이하로 선택해 주세요.")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("압축을 푼 양식이 너무 큽니다.")
        except zipfile.BadZipFile as exc:
            raise ValueError("정상적인 Excel/Word 파일을 선택해 주세요.") from exc
        if suffix == ".xlsx" and hashlib.sha256(content).digest() == hashlib.sha256(self.asset.read_bytes()).digest():
            result = copy.deepcopy(self.default)
            result.update(source_file=Path(filename).name, preview_only=True)
            return result
        names: list[tuple[str, str]] = []
        if suffix == ".xlsx":
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True, keep_links=False)
            try:
                for sheet in workbook:
                    if sheet.max_row > 2000 or sheet.max_column > 100:
                        raise ValueError("양식은 시트당 2,000행·100열 이하로 제한합니다.")
                    for row in sheet.iter_rows():
                        for index, cell in enumerate(row):
                            if not isinstance(cell.value, str) or cell.data_type == "f":
                                continue
                            label = cell.value.strip().rstrip(":：")
                            if not label or len(label) > 100 or index + 1 >= len(row):
                                continue
                            following = row[index + 1]
                            rgb = following.fill.fgColor.rgb if following.fill.fgColor.type == "rgb" else ""
                            if following.value is None and rgb in {TEAL, YELLOW}:
                                names.append((label, "auto" if rgb == TEAL else "user"))
                            elif index == 0 and (following.value is None or str(following.value).strip() in {"", "(미입력)", "입력"}):
                                names.append((label, "user"))
            finally:
                workbook.close()
        else:
            from docx import Document
            document = Document(io.BytesIO(content))
            for table in document.tables:
                for row in table.rows:
                    text = row.cells[0].text.strip().rstrip(":：")
                    if text and len(text) <= 100:
                        names.append((text, "user"))
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if ":" in text or "：" in text:
                    label = re.split("[:：]", text, maxsplit=1)[0].strip()
                    if label and len(label) <= 100:
                        names.append((label, "user"))
        fields, seen = [], set()
        known = {f["name"]: f for f in self.default["fields"]}
        for name, source in names:
            name = _canonical(name)
            if _key(name) in seen or name in {"항목", "값", "구분", "내용"}:
                continue
            seen.add(_key(name))
            mapping = known.get(name, {}).get("maps_to")
            if _blank_field(name):
                source = "blank"
            elif source == "auto" and not mapping:
                source = "user"
            fields.append(_field(name, source, **({"maps_to": mapping} if mapping and source == "auto" else {})))
        if not fields:
            raise ValueError("항목을 찾지 못했습니다. 항목 이름을 직접 등록해 주세요.")
        return {"name": "업로드 양식 미리보기", "is_default": False, "fields": fields[:200], "layout": "table", "source_file": Path(filename).name, "preview_only": True, "note": "항목과 입력 방식을 확인 후 저장하세요. 다운로드는 간단한 항목 표이며 원래 디자인을 복제하지 않습니다."}

    @staticmethod
    def _measurements(raw: Any, date: str) -> dict:
        if isinstance(raw, list):
            raw = {x.get("item", ""): x for x in raw if isinstance(x, dict)}
        raw = raw if isinstance(raw, dict) else {}
        if isinstance(raw.get("items"), list):
            raw = {x.get("item", ""): x for x in raw["items"] if isinstance(x, dict)}
        result = {}
        for item, _label, column, limit, unit in METRICS:
            entry = raw.get(item, raw.get(column, {}))
            if not isinstance(entry, dict):
                entry = {"value": entry}
            value = _number(entry.get("representative_value", entry.get("value")))
            stamp = entry.get("representative_timestamp") or entry.get("timestamp") or entry.get("measured_at")
            basis = entry.get("representative_basis", entry.get("basis", "측정값"))
            summary = entry.get("summary", {})
            if value is None and isinstance(summary, dict):
                value = _number(summary.get("minimum_observed"))
                if value is not None:
                    stamp = summary.get("minimum_observed_timestamp")
                    basis = "관측된 5초 평균 센서값 중 하루 최저값"
            if value is None:
                value = _number(entry.get("min"))
                basis = entry.get("basis", "해당 날짜의 최소값")
                stamp = stamp or entry.get("min_timestamp") or entry.get("min_time")
                hour = entry.get("min_hour")
                if stamp is None and hour is not None:
                    try:
                        stamp = f"{date} {int(hour):02d}:00 (시간 집계)"
                    except (TypeError, ValueError):
                        pass
            if value is None and isinstance(entry.get("series"), list):
                candidates = [(part, _number(part.get("min"))) for part in entry["series"] if isinstance(part, dict)]
                candidates = [(part, minimum) for part, minimum in candidates if minimum is not None]
                if candidates:
                    part, value = min(candidates, key=lambda pair: pair[1])
                    hour = part.get("timestamp") or part.get("hour")
                    stamp = f"{hour} (해당 시간구간 최소값)" if hour is not None else None
                    basis = "시간별 최소값 중 하루 최저값; 정확한 발생시각은 별도 원본 확인"
            if entry.get("data_available") is False:
                value, stamp = None, None
            threshold = _number(entry.get("threshold"))
            threshold = threshold if threshold is not None else float(limit)
            result[item] = {"value": value, "timestamp": stamp if value is not None else None, "threshold": threshold, "unit": entry.get("unit", unit), "basis": basis, "judgement": None if value is None else ("부적합" if value < threshold else "적합")}
        return result

    @staticmethod
    def _power(raw: Any) -> dict:
        raw = raw if isinstance(raw, dict) else {}
        kwh = _number(raw.get("total_kwh"))
        # Only physical kWh is accepted. Challenge-native Wh is not a report tariff input.
        unit = str(raw.get("energy_unit", raw.get("unit", "kWh"))).lower()
        if unit not in {"kwh", "physical_kwh"} or raw.get("data_available") is False:
            kwh = None
        cost = kwh * 180 if kwh is not None else None
        carbon = kwh * 0.424 if kwh is not None else None
        return {"total_kwh": kwh, "total_cost_krw": cost, "total_co2_kg": carbon, "max_demand_kw": _number(raw.get("max_demand_kw")), "equipment": raw.get("equipment"), "scope": raw.get("scope")}

    def create_report(self, date: str, equipment: str, action_taken: str = "", extra_fields: dict | None = None, measurements: Any = None, power: dict | None = None, anomalies: Any = None) -> dict:
        template = self.get_report_template()
        extras = {_key(_canonical(str(k))): v for k, v in (extra_fields or {}).items()}
        values = {"설비명": equipment, "조치 내용": action_taken}
        for field in template["fields"]:
            canonical = _canonical(field["name"])
            if _key(canonical) in extras:
                values[canonical] = extras[_key(canonical)]
        # A selected day is not a claim that an inspection took place at midnight.
        values.setdefault("점검 일시", None)
        action_value = values.get("조치 내용")
        no_action = _key(action_value) in {"", "none", "없음", "조치없음", "해당없음", "미조치", "미실시", "null"}
        action_required = (not no_action and not _missing(action_value)) or (extra_fields or {}).get("has_action") is True
        measured = self._measurements(measurements, date)
        current = self._power(power)
        previous_raw = (power or {}).get("previous", (power or {}).get("previous_power", {}))
        previous = self._power(previous_raw)
        if previous.get("equipment") and previous["equipment"] != current.get("equipment"):
            previous = self._power({})
        changes = {key: (current[key] / previous[key] - 1 if current[key] is not None and previous[key] not in (None, 0) else None) for key in ("total_kwh", "total_cost_krw", "total_co2_kg", "max_demand_kw")}
        factory = _key(current.get("equipment") or equipment) in {"전체", "공장전체", "factory", "all"} or current.get("scope") == "factory"
        contract = _number(values.get("계약전력 (kW)"))
        contract_ratio = current["max_demand_kw"] / contract if factory and current["max_demand_kw"] is not None and contract is not None and contract > 0 else None
        anomaly_input = anomalies if isinstance(anomalies, dict) else {"items": anomalies or []}
        events = [copy.deepcopy(x) for x in anomaly_input.get("items", []) if isinstance(x, dict)]
        ai_events = [copy.deepcopy(x) for x in anomaly_input.get("ai_items", []) if isinstance(x, dict)]
        available_anomalies = anomalies is not None and not anomaly_input.get("error") and anomaly_input.get("data_available") is not False
        context = {"equipment": equipment, "measurements": measured, "power": current, "previous_power": previous, "changes": changes, "contract_ratio": contract_ratio, "nonconforming_count": sum(v["judgement"] == "부적합" for v in measured.values()) if any(v["value"] is not None for v in measured.values()) else None, "anomaly_summary": f"조회 결과 {len(events)}건" if available_anomalies else None}
        rows = []
        for field in template["fields"]:
            row = copy.deepcopy(field)
            name = _canonical(field["name"])
            if field["source"] == "blank" or _blank_field(name):
                row.update(value=None, required=False, source="blank")
            elif field["source"] == "auto":
                row["value"] = _lookup(context, field.get("maps_to", ""))
            else:
                row["value"] = values.get(name)
            if action_required and name in ACTION:
                row["required"] = True
            rows.append(row)
        if action_required:
            for name in ACTION:
                if not any(_canonical(r["name"]) == name for r in rows):
                    rows.append({**_field(name, "user", required=True), "value": values.get(name, extras.get(_key(name)))})
        missing = [r["name"] for r in rows if r.get("required") and _missing(r.get("value"))]
        missing_auto = [r["name"] for r in rows if r.get("source") == "auto" and _missing(r.get("value"))]
        notes = ["실측 전력량(kWh) × 180원 및 × 0.424 kgCO₂로 계산합니다. 과정용 reference-native 값은 사용하지 않습니다.", "판정은 프로젝트 점검 임계값과의 비교이며 법정 안전 인증을 의미하지 않습니다."]
        if not factory:
            notes.append("개별 설비 전력과 공장 계약전력을 비교하지 않으므로 계약전력 대비는 비워 둡니다.")
        if (power or {}).get("peak_kw") is not None and current["max_demand_kw"] is None:
            notes.append("5초 평균 유효전력의 최댓값은 요금 산정용 최대 수요전력과 다르므로 최대 수요전력 칸에 대입하지 않습니다.")
        if not any(v["value"] is not None for v in measured.values()):
            notes.append("선택한 날짜의 센서 측정값이 없습니다. 측정·판정 칸을 비워 두며 부적합 건수 수식의 0은 점검 완료를 뜻하지 않습니다.")
        if any(("최소값" in v["basis"] or "최저값" in v["basis"]) and v["value"] is not None for v in measured.values()):
            notes.append("측정값은 해당 날짜의 최소값입니다. 시간 집계라고 표시된 시각은 해당 시간대를 가리키며, 각 항목의 최소 시각은 서로 다를 수 있습니다.")
        if action_required and not any(r["name"] == "조치 내용" and not _missing(r.get("value")) for r in rows):
            notes.append("조치가 있었다고 선택했습니다. 실제 조치 내용과 확인 결과를 입력해 주세요.")
        sources = []
        input_sources = [("measurements", measurements), ("power", power), ("anomalies", anomalies)]
        if isinstance(measurements, dict):
            input_sources.extend((f"measurements.{key}", value) for key, value in measurements.items() if isinstance(value, dict))
        for kind, raw in input_sources:
            if not isinstance(raw, dict):
                continue
            descriptors = {k: raw[k] for k in ("source", "source_file", "data_id", "source_id", "dataset_id", "source_period", "date", "data_period", "period_start", "period_end", "note", "model", "count", "ai_count", "truncated", "basis", "aggregation", "representative_basis") if k in raw}
            if descriptors:
                sources.append({"kind": kind, **descriptors})
            if raw.get("note"):
                notes.append(str(raw["note"]))
        # Explicitly retain the report query and common source identifier if provided.
        source_ids = {str(s[k]) for s in sources for k in ("data_id", "source_id", "dataset_id") if s.get(k)}
        if len(source_ids) > 1:
            notes.append("자료 식별자가 서로 다릅니다. 동일 원본과 기간인지 확인한 뒤 내보내세요.")
        return {"template": template, "name": template["name"], "date": date, "equipment": equipment, "rows": rows, "missing_user_fields": missing, "missing_auto_fields": missing_auto, "ready_to_export": not missing and len(source_ids) <= 1, "action_required": action_required, "events": events, "ai_events": ai_events, "measurements": measured, "power": current, "context": context, "sources": sources, "notes": list(dict.fromkeys(notes)), "data_id": next(iter(source_ids)) if len(source_ids) == 1 else None, "source_period": (power or {}).get("source_period") or anomaly_input.get("source_period"), "source_identity_mismatch": len(source_ids) > 1, "persisted": False}

    def export_xlsx(self, draft: dict) -> bytes:
        if draft.get("source_identity_mismatch"):
            raise ValueError("자료 식별자가 다릅니다. 같은 원본으로 보고서를 다시 생성해 주세요.")
        rows = draft.get("rows", [])
        missing = [r.get("name", "항목") for r in rows if r.get("required") and _missing(r.get("value"))]
        # Never trust only the caller-controlled ready_to_export flag.
        present = {_canonical(r.get("name", "")): r.get("value") for r in rows}
        missing.extend(name for name in CORE if _missing(present.get(name)) and name not in missing)
        if draft.get("action_required"):
            missing.extend(name for name in ACTION if _missing(present.get(name)) and name not in missing)
        if missing:
            raise ValueError("필수 항목을 입력해 주세요: " + ", ".join(missing))
        template = draft.get("template", {})
        if template.get("is_default"):
            workbook = openpyxl.load_workbook(self.asset, keep_links=False)
            sheet = workbook["안전점검보고서"]
            for row in rows:
                if row.get("cell") and not row.get("formula"):
                    sheet[row["cell"]] = None if row.get("source") == "blank" or _blank_field(row["name"]) else _safe_text(row.get("value"))
            # Color-only parsing cannot override these explicit blank requirements.
            for address in ("G1", "G2", "H2", "I2", "D61", "F61", "H61", "D62", "F62", "H62"):
                sheet[address] = None
            for i, (key, _name, _raw, limit, unit) in enumerate(METRICS, 17):
                measurement = draft.get("measurements", {}).get(key, {})
                threshold = measurement.get("threshold", limit)
                sheet[f"C{i}"] = f"{threshold:g} {unit} 이상"
                sheet[f"D{i}"].number_format = "0.00"
                sheet[f"E{i}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                sheet.row_dimensions[i].height = max(sheet.row_dimensions[i].height or 18, 30)
            for col in ("C", "D"):
                sheet[f"{col}31"].number_format = "#,##0.000"
                sheet[f"{col}32"].number_format = "#,##0"
                sheet[f"{col}33"].number_format = "#,##0.000"
                sheet[f"{col}34"].number_format = "#,##0.00"
            # Keep the ratio formula but provide no individual-equipment numerator.
            if draft.get("context", {}).get("contract_ratio") is None:
                sheet["C35"] = '=IF(AND(ISNUMBER(C34),ISNUMBER(C29),C29>0,F35="공장 전체 범위 확인"),C34/C29,"")'
                sheet["F35"] = "공장 전체·동일 범위 수요전력 자료가 있을 때만 비교"
            else:
                sheet["F35"] = "공장 전체 범위 확인"
            sheet["C26"] = '=IF(COUNTA(F17:F25)=0,"",COUNTIF(F17:F25,"부적합"))'
            sheet["F38"] = "편차 (%)"
            sheet["G38"] = "지속 시간 (분)"
            for rowno, event in zip(range(39, 43), draft.get("events", [])):
                self._write_event(sheet, rowno, event)
                sheet[f"B{rowno}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                sheet.row_dimensions[rowno].height = 30
                for col in ("D", "E", "F", "G"):
                    sheet[f"{col}{rowno}"].number_format = "0.00"
            if len(draft.get("events", [])) > 4:
                event_sheet = workbook.create_sheet("이상 내역 별지")
                self._event_sheet(event_sheet, draft)
        else:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "점검 보고서"
            sheet.append([draft.get("name", "점검 보고서"), "", "", ""])
            sheet.merge_cells("A1:D1")
            sheet.append(["사용자 항목 표 · 원본 디자인 복제 아님", "", "", ""])
            sheet.merge_cells("A2:D2")
            sheet.append(["항목", "내용", "입력 방식", "필수"])
            for row in rows:
                value = None if row.get("source") == "blank" or _blank_field(row["name"]) else row.get("value")
                sheet.append([_safe_text(row["name"]), _safe_text(value), {"auto": "자동", "user": "사용자", "blank": "공란"}.get(row.get("source"), ""), "필수" if row.get("required") else ""])
            self._style_table(sheet, {"A": 35, "B": 65, "C": 12, "D": 10})
            if draft.get("events"):
                self._event_sheet(workbook.create_sheet("이상 내역"), draft)
        if draft.get("ai_events"):
            ai_sheet = workbook.create_sheet("AI 탐지 구간")
            self._event_sheet(ai_sheet, {**draft, "events": draft["ai_events"]})
            ai_sheet["C2"] = "AI 이상 점수"
            ai_sheet["D2"] = "모델 기준 점수"
            ai_sheet["G1"] = "AI 탐지는 고장 원인 확정이 아닙니다. 센서 규칙 판정과 별도 결과입니다."
        self._source_sheet(workbook.create_sheet("자료 및 작성 기준"), draft)
        workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    @staticmethod
    def _write_event(sheet: Any, row: int, event: dict, offset: int = 2) -> None:
        for i, name in enumerate(("timestamp", "item", "value", "threshold", "deviation_pct", "duration_min"), offset):
            value = event.get(name)
            if name == "item":
                value = next((label for _key_, label, raw, _limit, _unit in METRICS if value in {_key_, raw}), value)
            if name == "duration_min" and value is None and _number(event.get("duration_seconds")) is not None:
                value = float(event["duration_seconds"]) / 60
            sheet.cell(row, i, _safe_text(value))
        if event.get("note"):
            sheet.cell(row, offset + 6, _safe_text(event["note"]))

    @classmethod
    def _event_sheet(cls, sheet: Any, draft: dict) -> None:
        sheet.append(["선택 날짜", draft.get("date"), "설비", draft.get("equipment")])
        sheet.append(["발생 일시", "이상 항목", "측정값", "기준값", "편차 (%)", "지속 시간 (분)", "비고"])
        for row, event in enumerate(draft.get("events", []), 3):
            cls._write_event(sheet, row, event, offset=1)
        cls._style_table(sheet, {"A": 25, "B": 24, "C": 16, "D": 16, "E": 14, "F": 18, "G": 48})
        for cell in sheet[2]:
            cell.fill = PatternFill("solid", fgColor="DDEBE8")
            cell.font = Font(name="맑은 고딕", bold=True, size=10)
        sheet.print_title_rows = "1:2"
        for row in sheet.iter_rows(min_row=3, min_col=3, max_col=6):
            for cell in row:
                cell.number_format = "0.000"

    @classmethod
    def _source_sheet(cls, sheet: Any, draft: dict) -> None:
        sheet.append(["항목", "내용"])
        sheet.append(["조회 날짜", draft.get("date")])
        sheet.append(["설비", _safe_text(draft.get("equipment"))])
        for key in ("data_id", "source_period"):
            if draft.get(key) is not None:
                sheet.append([key, _safe_text(draft[key])])
        for source in draft.get("sources", []):
            sheet.append([source.get("kind", "자료"), _safe_text(source)])
        for note in draft.get("notes", []):
            sheet.append(["작성 기준", _safe_text(note)])
        sheet.append(["필수 입력", "점검 일시·설비명·점검자·종합 판정. 조치가 있으면 조치 내용·일시·조치자·조치 후 확인."])
        sheet.append(["계산", "수식은 Excel/호환 프로그램에서 파일을 열 때 다시 계산합니다."])
        cls._style_table(sheet, {"A": 22, "B": 100})

    @staticmethod
    def _style_table(sheet: Any, widths: dict) -> None:
        sheet.freeze_panes = "A3"
        sheet.sheet_view.showGridLines = False
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        for row in sheet:
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="DDEBE8")
            cell.font = Font(name="맑은 고딕", bold=True, size=11)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
