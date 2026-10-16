"""Notification policy remains separate from original detections and report decisions."""
import copy
import io
import openpyxl
import pytest
from agent import agent_tools as mod
from test_report_flow import tools, fill_required, make_events, summary


def snapshot(available=True):
    return {"status": "ok" if available else "unavailable", "data_available": available,
        "policy": {"score_margin": .02, "duration_seconds": 60, "repeat_count": 3,
            "repeat_window_minutes": 30, "merge_window_minutes": 10},
        "summary": {"notification_count": 1} if available else {},
        "coverage": {}, "note": "local preview only", "notifications": [{
            "equipment": "1 (PM-3)", "first_trigger_at": "2025-04-28T01:20:05",
            "last_trigger_at": "2025-04-28T01:25:05", "first_score_at_trigger": .7,
            "score_max_at_triggers": .72, "candidate_count": 2, "reasons": ["repeated"],
            "source_event_ids": ["e1", "e2", "e3"]}] if available else []}


def set_evidence(monkeypatch, tools):
    ai = [{**x, "kind": "ai", "item": "AI", "item_key": "ai", "value": .7,
           "unit": "anomaly_score", "threshold": .6} for x in make_events(3)]
    s = summary(make_events())
    s["ai_events"] = ai
    s["ai"].update(event_count=3, flagged_records=3)
    monkeypatch.setattr(mod, "alert_summary", lambda *a, **k: copy.deepcopy(s))
    monkeypatch.setattr(mod, "ai_notification_snapshot", lambda *a, **k: snapshot())
    monkeypatch.setattr(mod, "sensor_insight_snapshot", lambda *a, **k: {"data_available": False, "summary": {}, "policy": {}, "events": [], "groups": []})
    return ai


def make_legacy_draft(tools, available=True):
    """An already-open pre-sensor draft must keep its original policy evidence."""
    tools.create_report()
    tools.state["report_draft"].pop("sensor_insight_snapshot", None)
    tools.state["report_draft"]["ai_notification_snapshot"] = snapshot(available)


def test_agent_retains_ai_count_and_adds_separate_notifications(monkeypatch, tools):
    ai = set_evidence(monkeypatch, tools)
    result = tools.get_anomalies()
    assert result["ai_count"] == 3
    assert result["ai_items"] == ai
    assert result["ai_notifications"]["summary"]["notification_count"] == 1
    assert result["ai_notifications"]["notifications"][0]["first_trigger_at"] != ai[0]["timestamp"]


def test_report_snapshot_pinned_during_policy_change(monkeypatch, tools):
    set_evidence(monkeypatch, tools)
    make_legacy_draft(tools)
    saved = copy.deepcopy(tools.state["report_draft"]["ai_notification_snapshot"])
    tools.state["ai_notification_policy"] = {"score_margin": .05}
    monkeypatch.setattr(mod, "ai_notification_snapshot", lambda *a, **k: pytest.fail("Existing draft snapshot must remain pinned"))
    tools.create_report(extra_fields={"점검자": "테스트 담당자"})
    assert tools.state["report_draft"]["ai_notification_snapshot"] == saved
    assert len(tools.state["report_draft"]["ai_events"]) == 3


def test_export_keeps_full_ai_sheet_and_separate_notification_sheet(monkeypatch, tools):
    ai = set_evidence(monkeypatch, tools)
    make_legacy_draft(tools)
    fill_required(tools)
    tools.confirm_judgment("추가 점검 필요")
    draft = tools.state["report_draft"]
    assert draft["ai_events"] == ai
    content = mod.export_xlsx(draft, tools.templates.get())
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb["AI 탐지 전체 내역"].max_row == len(ai) + 1
    assert wb["AI 알림 후보"].max_row == 2
    assert wb["AI 알림 후보"]["B2"].value == "2025-04-28T01:20:05"
    assert wb["AI 알림 후보"]["D2"].value == .7
    assert any("연속 이상 시간" in str(row[1].value) for row in wb["집계 기준 및 확인"].iter_rows())


def test_missing_notification_cache_is_unknown_in_export(monkeypatch, tools):
    monkeypatch.setattr(mod, "sensor_insight_snapshot", lambda *a, **k: {"data_available": False, "summary": {}, "events": [], "groups": []})
    monkeypatch.setattr(mod, "ai_notification_snapshot", lambda *a, **k: snapshot(False))
    make_legacy_draft(tools, available=False)
    fill_required(tools)
    tools.confirm_judgment("추가 점검 필요")
    wb = openpyxl.load_workbook(io.BytesIO(mod.export_xlsx(tools.state["report_draft"], tools.templates.get())))
    rows = {row[0].value: row[1].value for row in wb["집계 기준 및 확인"].iter_rows()}
    assert rows["알림 후보 묶음"] == "확인 불가"
    assert wb["AI 알림 후보"].max_row == 1
