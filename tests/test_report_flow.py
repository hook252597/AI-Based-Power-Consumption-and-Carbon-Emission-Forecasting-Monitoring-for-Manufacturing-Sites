import copy
import io
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from agent import agent_tools as mod
from agent.agent_core import opening_summary, demo_reply, state_block
from agent.agent_tools import AgentTools, TemplateStore, duration_text, export_xlsx

EQ = '1 (PM-3)'


def make_events(n=7, seconds=5):
    return [{'equipment': EQ, 'item': '상전압 R', 'item_key': 'phase_voltage_R',
             'timestamp': f'2025-04-28T01:{i:02d}:00', 'end': f'2025-04-28T01:{i:02d}:05',
             'value': 205.0, 'unit': 'V', 'threshold': 210.0,
             'flagged_records': 1, 'duration_seconds': seconds,
             'duration_min': seconds / 60 if seconds is not None else None, 'basis': 'exact_5s'} for i in range(n)]


def summary(events, exact=True):
    rule = dict(equipment_count=1, event_count=len(events) if exact else None,
                affected_hours=1, flagged_records=len(events) if exact else None,
                sample_duration_seconds=len(events)*5 if exact else None, longest_duration_seconds=5 if exact else None)
    ai = dict(equipment_count=0,event_count=0 if exact else None,affected_hours=0,flagged_records=0,sample_duration_seconds=0,longest_duration_seconds=0 if exact else None)
    return dict(date='2025-04-28', data_available=True, basis='exact_5s' if exact else 'hourly',
        rule=rule, ai=ai, equipment_rows=[dict(equipment=EQ,data_available=True,rule=rule,ai=ai)],
        rule_events=events, ai_events=[], note='test source')


@pytest.fixture
def tools(monkeypatch):
    frame = pd.DataFrame([{'hour':1, 'kwh':3.2, 'equipment': EQ,
         **{col: 205 if fam == 'phase_voltage' else 370 if fam == 'line_voltage' else 90
            for col, label, unit, fam in mod.ITEMS.values()}}])
    monkeypatch.setattr(mod, '_day', lambda d, equipment=None: frame.copy())
    monkeypatch.setattr(mod, 'alert_summary', lambda d, thresholds, equipment=None: summary(make_events()))
    state = {'view_date':'2025-04-28','selected_equipment':EQ}
    store = TemplateStore(state=state)
    return AgentTools(state, store)


def fill_required(tools):
    tools.create_report()
    values = {r['name']:'테스트 담당자' if r['name']=='점검자' else '정기 점검'
              for r in tools.state['report_draft']['rows'] if r.get('source')=='user' and r.get('required')}
    values['종합판정'] = '적합'
    tools.create_report(extra_fields=values)


def test_session_templates_never_write_shared_files(tmp_path):
    original = mod.APP_DIR / mod.DEFAULT_TEMPLATE_FILE
    (tmp_path / mod.DEFAULT_TEMPLATE_FILE).write_bytes(original.read_bytes())
    a, b = {}, {}
    sa, sb = TemplateStore(tmp_path,a), TemplateStore(tmp_path,b)
    sa.ensure_default(); sb.ensure_default()
    sa.save([dict(name='작성자',source='user',required=True)])
    assert not sa.get()['is_default']
    assert sb.get()['is_default']
    assert not sb.template_asked()
    assert sorted(p.name for p in tmp_path.iterdir()) == [mod.DEFAULT_TEMPLATE_FILE]


def test_uploaded_template_kept_only_in_own_session(tmp_path):
    sa, sb = TemplateStore(tmp_path,{}), TemplateStore(tmp_path,{})
    content=(mod.APP_DIR/mod.DEFAULT_TEMPLATE_FILE).read_bytes()
    token=sa.store_upload('company.xlsx',content)
    saved=sa.save([dict(name='설비명',source='auto')],layout='xlsx',source_path=str(token),source_file='company.xlsx')
    assert saved['_source_bytes']==content
    assert not sb.data.get('uploads')
    assert list(tmp_path.iterdir())==[]


def test_model_cannot_confirm_judgment_and_export_guard(tools):
    fill_required(tools)
    draft=tools.state['report_draft']
    assert next(v for k,v in draft['user_values'].items() if mod._norm(k)=='종합판정')=='적합'
    assert not draft['judgment_confirmed']
    assert not draft['ready_to_export']
    assert tools.execute('confirm_judgment',{'judgment':'적합'}).get('error')
    with pytest.raises(ValueError):
        export_xlsx(draft,tools.templates.get())
    tools.confirm_judgment('적합')
    assert draft['ready_to_export']


def test_field_can_be_cleared_and_confirmation_invalidated(tools):
    fill_required(tools)
    tools.confirm_judgment('적합')
    result=tools.create_report(extra_fields={'점검자':''})
    assert '점검자' in result['missing_user_fields']
    assert '점검자' not in result['user_values']
    assert not result['judgment_confirmed']
    assert not result['ready_to_export']


def test_action_can_be_cleared(tools):
    result=tools.create_report(action_taken='점검 완료')
    key=next(n for n in result['user_values'] if '조치내용' in mod._norm(n))
    result=tools.create_report(extra_fields={key:''})
    assert result['action_taken']==''
    assert key not in result['user_values']


def test_target_changes_keep_existing_draft(tools):
    tools.create_report(extra_fields={'점검자':'김 담당'})
    before=copy.deepcopy(tools.state['report_draft'])
    tools.state.update(view_date='2025-04-29',selected_equipment='12 (4호기)')
    result=tools.create_report(date='2025-04-29',equipment='4호기')
    assert result['target_conflict']
    assert tools.state['report_draft']==before
    tools.start_new_report('2025-04-29','4호기')
    assert tools.state['previous_report_draft']==before
    assert tools.state['report_draft']['equipment']=='12 (4호기)'
    assert tools.state['report_draft']['user_values']=={}


def test_omitted_target_updates_existing_draft(tools):
    tools.create_report()
    tools.state.update(view_date='2025-04-29',selected_equipment='12 (4호기)')
    result=tools.create_report(extra_fields={'점검자':'김 담당'})
    assert result['date']=='2025-04-28'
    assert result['equipment']==EQ


def test_all_events_exported_and_five_seconds_remains_five_seconds(tools):
    fill_required(tools)
    tools.confirm_judgment('추가 점검 필요')
    draft=tools.state['report_draft']
    tables=[r for r in draft['rows'] if r.get('kind')=='table']
    assert tables and len(tables[0]['value'])==7
    assert tables[0]['value'][0]['지속 시간']=='5초'
    data=export_xlsx(draft,tools.templates.get())
    wb=openpyxl.load_workbook(io.BytesIO(data))
    assert wb['설정 기준 전체 내역'].max_row==8
    assert wb['설정 기준 전체 내역']['I2'].value=='5초'
    assert wb['AI 탐지 전체 내역'].max_row==1
    notes=' '.join(str(c.value) for row in wb['집계 기준 및 확인'] for c in row)
    assert '전체 내역 7개' in notes
    assert '추가 점검 필요' in notes


def test_contract_power_and_max_demand_not_invented(tools):
    result=tools.create_report()
    values={row['name']:row.get('value') for row in result['rows'] if row.get('source')=='auto'}
    selected=[v for k,v in values.items() if '최대수요전력' in mod._norm(k) or '계약전력' in mod._norm(k)]
    assert selected
    assert all(v is None for v in selected)


def test_hourly_fallback_duration_and_count_not_fabricated(tools,monkeypatch):
    fallback=summary(make_events(1,seconds=None),exact=False)
    monkeypatch.setattr(mod,'alert_summary',lambda *args,**kwargs:fallback)
    result=tools.get_anomalies()
    assert result['count'] is None
    assert result['items'][0]['duration_seconds'] is None
    assert duration_text(result['items'][0])=='확인 불가 · 이상 기록이 관측된 시간대'
    text=opening_summary(tools.state,tools)
    assert '구간 수 확인 불가' in text
    assert '1시간' not in text
    assert 'None' not in text


def test_missing_data_has_no_normal_claim(tools,monkeypatch):
    missing=summary([])
    missing.update(data_available=False,equipment_rows=[])
    monkeypatch.setattr(mod,'alert_summary',lambda *args,**kwargs:missing)
    text=opening_summary(tools.state,tools)
    assert '판단할 수 없습니다' in text
    assert '정상' not in text


def test_threshold_comparison_respects_unknown_event_count(tools,monkeypatch):
    monkeypatch.setattr(mod,'alert_summary',lambda *args,**kwargs:summary(make_events(1,None),False))
    result=tools.set_thresholds(phase_voltage=209)
    assert result['anomaly_count_after'] is None
    assert result['anomaly_count_before'] is None


def test_template_change_requires_new_judgment(tools):
    fill_required(tools); tools.confirm_judgment('적합')
    tools.save_report_template([{'name':'설비명','source':'auto'},{'name':'종합판정','source':'auto'}])
    assert tools.templates.get()['fields'][1]['source']=='user'
    assert not tools.state['report_draft']['judgment_confirmed']


def test_unknown_fields_do_not_mutate_draft(tools):
    tools.create_report()
    old=copy.deepcopy(tools.state['report_draft'])
    assert tools.create_report(extra_fields={'존재하지 않는 항목':'value'}).get('error')
    assert tools.state['report_draft']==old


def test_state_prompt_marks_demo_and_no_fake_durations(tools,monkeypatch):
    import agent.agent_core as core
    monkeypatch.setattr(core,'sensor_period',lambda:('2024-12-01','2025-04-30 00:00'))
    monkeypatch.setattr(core,'forecast_period',lambda:('2025-04-30','2025-05-31 00:00'))
    text=state_block(tools.state,tools)
    assert '실제 오늘 날짜와 다름' in text
    assert '시간대 수를 시간 단위 지속 시간으로 바꾸지 않는다' in text
    assert '지속시간은' not in text

def test_report_ui_guided_flow_without_model_calls(monkeypatch):
    from streamlit.testing.v1 import AppTest
    import agent.agent_ui as ui
    monkeypatch.setattr(ui,'load_settings',lambda:('', 'gpt-5-mini'))
    script = """
import streamlit as st
from agent.agent_ui import render_agent_tab
st.session_state.setdefault('view_date','2025-04-28')
st.session_state.setdefault('selected_equipment','1 (PM-3)')
render_agent_tab()
"""
    app=AppTest.from_string(script,default_timeout=30).run()
    assert not app.exception
    app.button(key='ag_start_report').click().run()
    assert not app.exception
    assert 'report_draft' in app.session_state
    assert any('조치 기록과 필수 정보' in (x.value or '') for x in app.markdown)
    for field in app.text_input:
        if field.label.endswith('*'):
            field.set_value('테스트 담당자' if '점검자' in field.label else '정기 점검')
    next(b for b in app.button if b.label=='입력 내용을 보고서에 반영').click().run()
    assert not app.exception
    app.selectbox[0].select('추가 점검 필요')
    app.button(key='ag_confirm_judgment').click().run()
    assert not app.exception
    assert app.session_state['report_draft']['ready_to_export']
    original=copy.deepcopy(app.session_state['report_draft'])
    app.session_state['view_date']='2025-04-29'
    app.session_state['selected_equipment']='12 (4호기)'
    app.run()
    assert not app.exception
    assert app.session_state['report_draft']==original
    assert any('작성 중: 2025-04-28 · 1 (PM-3) · 생성 당시 자료 고정' in x.value for x in app.info)



def test_report_handoff_starts_once_and_protects_existing_draft(monkeypatch):
    from streamlit.testing.v1 import AppTest
    import agent.agent_ui as ui
    monkeypatch.setattr(ui, 'load_settings', lambda: ('', 'gpt-5-mini'))
    script = """
import streamlit as st
from agent.agent_ui import render_agent_tab
st.session_state.setdefault('view_date', '2025-04-28')
st.session_state.setdefault('selected_equipment', '1 (PM-3)')
if not st.session_state.get('_test_initialized'):
    st.session_state['report_start_requested'] = True
    st.session_state['_test_initialized'] = True
render_agent_tab()
"""
    app = AppTest.from_string(script, default_timeout=30).run()
    assert not app.exception
    draft = copy.deepcopy(app.session_state['report_draft'])
    assert draft['date'] == '2025-04-28'
    assert draft['equipment'] == '1 (PM-3)'
    assert 'report_start_requested' not in app.session_state
    app.run()
    assert app.session_state['report_draft'] == draft
    app.session_state['view_date'] = '2025-04-29'
    app.session_state['selected_equipment'] = '12 (4호기)'
    app.session_state['report_start_requested'] = True
    app.run()
    assert not app.exception
    assert app.session_state['report_draft'] == draft
    assert any(x.label == 'AI에게 자료 질문하기' and not x.proto.expanded for x in app.expander)
    assert app.session_state['report_start_requested']
    assert app.button(key='ag_new_report')


def test_missing_secrets_does_not_show_internal_paths(monkeypatch, tmp_path):
    from streamlit.testing.v1 import AppTest
    import agent.agent_ui as ui
    monkeypatch.setattr(ui, '_configured_secret_files', lambda: [str(tmp_path / 'missing-secret.toml')])
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)
    monkeypatch.delenv('OPENAI_MODEL', raising=False)
    app = AppTest.from_string("from agent.agent_ui import load_settings\nimport streamlit as st\nkey, model = load_settings()\nst.caption('AI mode' if key else 'local mode')", default_timeout=20).run()
    assert not app.error
    assert not app.exception
    assert app.caption[0].value == 'local mode'


def test_environment_credentials_work_without_secret_file(monkeypatch, tmp_path):
    import agent.agent_ui as ui
    monkeypatch.setattr(ui, '_configured_secret_files', lambda: [str(tmp_path / 'missing-secret.toml')])
    monkeypatch.setenv('OPENAI_API_KEY', 'test-only-not-a-real-key')
    monkeypatch.setenv('OPENAI_MODEL', 'gpt-5-mini')
    key, model = ui.load_settings()
    assert key == 'test-only-not-a-real-key'
    assert model == 'gpt-5-mini'


def test_existing_secret_file_precedence_and_invalid_file_fallback(monkeypatch, tmp_path):
    import agent.agent_ui as ui
    path = tmp_path / 'settings.toml'
    path.write_text('OPENAI_API_KEY = "file-test-placeholder"\nOPENAI_MODEL = "gpt-5-mini"', encoding='utf-8')
    monkeypatch.setattr(ui, '_configured_secret_files', lambda: [str(path)])
    monkeypatch.setenv('OPENAI_API_KEY', 'env-test-placeholder')
    assert ui.load_settings()[0] == 'file-test-placeholder'
    path.write_text('invalid = [', encoding='utf-8')
    assert ui.load_settings()[0] == 'env-test-placeholder'
