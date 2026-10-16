"""Scoped sensor evidence is frozen in reports and separate from original AI logs."""
import copy
import io
import json
import openpyxl
import pytest
from agent import agent_tools as mod
from agent.agent_core import MAX_TOOL_CHARS, _compact, state_block
from monitoring import sensor_insights as backend
from test_report_flow import EQ, tools, fill_required, make_events, summary

OTHER='12 (4호기)'
POLICY=dict(backend.DEFAULT_POLICY)


def event(equipment=EQ,n=1):
    return dict(id=f's{n}',equipment=equipment,item='powerFactorR',kind='pf_sharp',direction='low',
        start='2025-04-28T01:00:00',end='2025-04-28T01:00:10',trigger_at='2025-04-28T01:00:05',
        value_at_trigger=65.,extreme_value=61.,duration_seconds=10,duration_at_trigger_seconds=5,
        baseline=94.,drop_pp=29.,ai_score=.7,trigger_ai_event_ids=[f'ai{n}'],source_ai_event_ids=[f'ai{n}',f'ai{n}_later'])


def group(equipment=EQ,n=1):
    return dict(id=f'g{n}',equipment=equipment,item='powerFactorR',kind='pf_sharp',direction='low',
        first_trigger_at='2025-04-28T01:00:05',last_trigger_at='2025-04-28T01:09:00',
        first_value_at_trigger=65.,event_count=2,event_ids=[f's{n}',f's{n}_later'],
        extreme_value=61.,max_event_duration_seconds=10,source_ai_event_ids=[f'ai{n}',f'ai{n}_later'])


def counts(n):
    return dict(original_event_count=n,ai_record_count=n,event_count=n,group_count=n,
        voltage_severe_count=0,voltage_sustained_count=0,pf_sharp_count=n,pf_sustained_count=0)


def source(day='2025-04-28',policy=None,available=True):
    per=[dict(equipment=EQ,data_available=available,observed_records=120,**counts(1)),
         dict(equipment=OTHER,data_available=available,observed_records=240,**counts(2))]
    if not available:
        for row in per: row.update(dict.fromkeys(mod.SENSOR_COUNT_KEYS))
    return dict(day=day,status='ok' if available else 'unavailable',data_available=available,policy=dict(policy or POLICY),
        summary=counts(3) if available else dict.fromkeys(mod.SENSOR_COUNT_KEYS),per_equipment=per,
        events=[event(),event(OTHER,2),event(OTHER,3)] if available else [],
        groups=[group(),group(OTHER,2),group(OTHER,3)] if available else [],
        coverage=dict(dates=[day],sample_seconds=5,observed_records=360),
        source=dict(partition_sha256='verified-partition',model_sha256='unchanged-model'),note='실험용 220V 가정')


@pytest.fixture(autouse=True)
def no_external_or_large_data(monkeypatch):
    monkeypatch.setattr(backend,'get_sensor_insights',lambda day,policy=None,cutoff=None:source(day,policy,day<'2025-04-30'))
    monkeypatch.setattr(mod,'ai_notification_snapshot',lambda *a,**k:dict(status='ok',data_available=True,
        summary={'notification_count':1},notifications=[],policy={},note='legacy'))


def selected_context():
    src=source()
    return dict(day='2025-04-28',date='2025-04-28',equipment=EQ,policy=POLICY.copy(),
        summary=copy.deepcopy(src['per_equipment'][0]),events=[event()],groups=[group()],
        source=src['source'],coverage=src['coverage'],data_available=True,status='ok')


def test_snapshot_filters_selected_equipment_and_uses_its_summary():
    snap=mod.sensor_insight_snapshot('2025-04-28','4호기')
    assert snap['equipment']==OTHER and snap['summary']['event_count']==2
    assert snap['coverage']['observed_records']==240
    assert len(snap['events'])==len(snap['groups'])==2
    assert all(x['equipment']==OTHER for x in snap['events']+snap['groups']+snap['per_equipment'])
    assert snap['source']['partition_sha256']=='verified-partition'


def test_missing_selected_equipment_does_not_use_global_zero(monkeypatch):
    result=source();result['per_equipment'][1]['data_available']=False
    monkeypatch.setattr(backend,'get_sensor_insights',lambda *a,**k:result)
    snap=mod.sensor_insight_snapshot('2025-04-28',OTHER)
    assert not snap['data_available'] and snap['summary']['event_count'] is None
    assert snap['events']==snap['groups']==[]
    with pytest.raises(ValueError): mod.sensor_insight_snapshot('2025-04-28','없는설비')


def test_clicked_report_context_is_consumed_and_frozen(monkeypatch,tools):
    ctx=selected_context();tools.state['sensor_insight_report_context']=ctx
    monkeypatch.setattr(backend,'get_sensor_insights',lambda *a,**k:pytest.fail('Clicked evidence must be used'))
    tools.create_report();snap=tools.state['report_draft']['sensor_insight_snapshot']
    assert snap['captured_from']=='dashboard_selection' and snap['summary']['event_count']==1
    assert snap['source']['model_sha256']=='unchanged-model'
    assert 'sensor_insight_report_context' not in tools.state
    ctx['events'][0]['extreme_value']=0;ctx['policy']['pf_low']=75
    assert snap['events'][0]['extreme_value']==61 and snap['policy']['pf_low']==80


@pytest.mark.parametrize('change',[{'day':'2025-04-27'},{'date':'2025-04-27'},{'equipment':OTHER}])
def test_stale_report_context_does_not_replace_target(change,tools):
    ctx=selected_context();ctx.update(change);tools.state['sensor_insight_report_context']=ctx
    tools.create_report();snap=tools.state['report_draft']['sensor_insight_snapshot']
    assert snap['captured_from']=='verified_sensor_cache' and snap['equipment']==EQ and snap['day']=='2025-04-28'
    assert 'sensor_insight_report_context' not in tools.state


def test_report_snapshot_frozen_when_policy_date_or_equipment_changes(monkeypatch,tools):
    tools.create_report();saved=copy.deepcopy(tools.state['report_draft']['sensor_insight_snapshot'])
    tools.state.update(view_date='2025-04-29',selected_equipment=OTHER,sensor_insight_policy={'pf_low':75})
    monkeypatch.setattr(mod,'sensor_insight_snapshot',lambda *a,**k:pytest.fail('Frozen evidence must not recompute'))
    result=tools.create_report(extra_fields={'점검자':'담당자'})
    assert result['date']=='2025-04-28' and result['equipment']==EQ
    assert tools.state['report_draft']['sensor_insight_snapshot']==saved
    before=copy.deepcopy(tools.state['report_draft'])
    assert tools.create_report(date='2025-04-29',equipment=OTHER)['target_conflict']
    assert tools.state['report_draft']==before


def test_explicit_new_report_captures_current_policy_and_target(tools):
    tools.create_report();before=copy.deepcopy(tools.state['report_draft'])
    tools.state['sensor_insight_policy']={**POLICY,'pf_low':75};tools.start_new_report('2025-04-29',OTHER)
    draft=tools.state['report_draft'];snap=draft['sensor_insight_snapshot']
    assert tools.state['previous_report_draft']==before
    assert snap['policy']['pf_low']==75 and snap['day']=='2025-04-29' and snap['equipment']==OTHER


def test_old_open_draft_is_not_silently_migrated(monkeypatch,tools):
    tools.create_report();draft=tools.state['report_draft'];draft.pop('sensor_insight_snapshot')
    draft['ai_notification_snapshot']={'policy':{'score_margin':.02},'data_available':True,'notifications':[],'summary':{'notification_count':0}}
    old=copy.deepcopy(draft['ai_notification_snapshot'])
    monkeypatch.setattr(mod,'sensor_insight_snapshot',lambda *a,**k:pytest.fail('Old draft must keep its policy'))
    tools.create_report(extra_fields={'점검자':'담당자'})
    assert tools.state['report_draft']['ai_notification_snapshot']==old
    assert not tools.state['report_draft']['sensor_insight_snapshot']


def test_tool_keeps_original_ai_counts_and_separate_sensor_scope(monkeypatch,tools):
    raw=summary(make_events());raw['ai']=dict(event_count=7);raw['ai_events']=[{**r,'kind':'ai'} for r in make_events()]
    monkeypatch.setattr(mod,'alert_summary',lambda *a,**k:copy.deepcopy(raw))
    out=tools.get_anomalies(equipment=OTHER)
    assert out['ai_count']==7 and out['ai_items']==raw['ai_events']
    assert out['sensor_insights']['summary']['event_count']==2 and out['ai_notifications']['legacy_compatibility']
    assert 'sensor_insights' in state_block(tools.state,tools)


def test_tool_copy_bounded_report_retains_every_event(monkeypatch,tools):
    big=mod.sensor_insight_snapshot('2025-04-28',EQ)
    big['events']=[{**event(n=n),'source_ai_event_ids':[f'ai{n}_{k}' for k in range(300)]} for n in range(100)]
    big['groups']=[group(n=n) for n in range(100)];big['summary'].update(event_count=100,group_count=100)
    monkeypatch.setattr(mod,'sensor_insight_snapshot',lambda *a,**k:copy.deepcopy(big))
    out=tools.create_report();snap=out['sensor_insight_snapshot']
    assert len(snap['events'])==12 and snap['total_event_count']==100 and len(snap['events'][0]['source_ai_event_ids'])==6
    stored=tools.state['report_draft']['sensor_insight_snapshot']
    assert len(stored['events'])==100 and len(stored['events'][0]['source_ai_event_ids'])==300
    text=_compact(out);small=json.loads(text)['sensor_insight_snapshot']
    assert len(text)<=MAX_TOOL_CHARS and small['total_event_count']==100 and small['summary']['event_count']==100
    assert len(stored['events'])==100


def test_export_preserves_ai_and_distinguishes_real_duration_from_group_span(tools):
    fill_required(tools);tools.confirm_judgment('추가 점검 필요')
    draft=tools.state['report_draft'];draft['ai_events']=[{**r,'kind':'ai'} for r in make_events(3)]
    wb=openpyxl.load_workbook(io.BytesIO(mod.export_xlsx(draft,tools.templates.get())))
    assert wb['AI 탐지 전체 내역'].max_row==4 and 'AI 알림 후보' not in wb.sheetnames
    assert wb['센서 점검 정보'].max_row==2
    assert wb['센서 점검 정보']['K2'].value==10 and wb['센서 점검 정보']['L2'].value==5
    assert wb['센서 점검 정보']['P2'].value=='ai1' and wb['센서 점검 정보']['Q2'].value=='ai1, ai1_later'
    assert wb['센서 정보 묶음']['K2'].value==10 and '구간 중 고정' in wb['센서 점검 정보']['M1'].value
    notes=' '.join(str(c.value) for row in wb['집계 기준 및 확인'] for c in row)
    assert '공백은 지속 시간에 더하지 않습니다' in notes and '220V' in notes


def test_missing_sensor_data_is_unknown_in_tool_and_export(tools):
    sensor=tools.get_anomalies(date='2025-05-01',equipment=EQ)['sensor_insights']
    assert not sensor['data_available'] and sensor['summary']['event_count'] is None and sensor['total_event_count'] is None
    tools.state['view_date']='2025-05-01';fill_required(tools);tools.confirm_judgment('추가 점검 필요')
    wb=openpyxl.load_workbook(io.BytesIO(mod.export_xlsx(tools.state['report_draft'],tools.templates.get())))
    evidence={r[0].value:r[1].value for r in wb['집계 기준 및 확인'].iter_rows()}
    assert evidence['센서 정보 구간 수']=='확인 불가' and evidence['센서 정보 묶음 수']=='확인 불가'
    assert wb['센서 점검 정보'].max_row==1


def test_policy_change_refreshes_opening_summary_without_external_call(monkeypatch,tools):
    from agent import agent_ui
    calls=[];monkeypatch.setattr(agent_ui,'opening_summary',lambda *a:calls.append(1) or 'local summary')
    agent_ui._ensure_summary(tools.state,tools);agent_ui._ensure_summary(tools.state,tools)
    assert len(calls)==1
    tools.state['sensor_insight_policy']={'pf_low':75};agent_ui._ensure_summary(tools.state,tools)
    assert len(calls)==2 and '조건을 바꾸어' in tools.state['messages'][-1]['content']


def test_pending_old_selection_does_not_override_new_policy(tools):
    tools.state['sensor_insight_report_context']=selected_context()
    tools.state['sensor_insight_policy']={**POLICY,'pf_low':75}
    tools.create_report()
    snap=tools.state['report_draft']['sensor_insight_snapshot']
    assert snap['captured_from']=='verified_sensor_cache'
    assert snap['policy']['pf_low']==75


def test_unavailable_selected_context_never_becomes_zero(tools):
    ctx=selected_context();ctx['status']='unavailable';ctx['data_available']=False
    tools.state['sensor_insight_report_context']=ctx
    tools.create_report();snap=tools.state['report_draft']['sensor_insight_snapshot']
    assert not snap['data_available'] and snap['summary']['event_count'] is None
    assert snap['events']==snap['groups']==[]
